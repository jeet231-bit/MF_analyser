"""Robustness: odd workbooks, concurrent users, and a new version arriving mid-session."""

import threading
import time
from concurrent.futures import ThreadPoolExecutor

import openpyxl
import pytest
from fastapi.testclient import TestClient

from app.config import get_settings
from app.main import create_app
from app.storage import runs as run_store
from tests.fixtures.logic_workbook import logic_fixture_xlsx_bytes
from tests.fixtures.workbook import inject_cached_values, workbook_bytes

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
SCOPE = ["Inputs", "Lookup", "Series", "Outputs", "Semantics"]


def _upload(client: TestClient, name: str, data: bytes) -> dict:
    response = client.post("/api/workbooks", files={"file": (name, data, XLSX)})
    assert response.status_code == 201, response.text
    return response.json()


def _odd_workbook() -> bytes:
    """An empty sheet, a merged header over a small table, and a sheet of only text."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Table"
    ws.merge_cells("A1:C1")
    ws["A1"] = "Quarterly sales (merged title)"
    ws["A2"], ws["B2"], ws["C2"] = "Region", "Sales", "Share"
    total = 0
    for i, (region, sales) in enumerate([("North", 10), ("South", 30), ("East", 60)], start=3):
        ws[f"A{i}"], ws[f"B{i}"], ws[f"C{i}"] = region, sales, f"=B{i}/SUM($B$3:$B$5)"
        total += sales
    wb.create_sheet("Empty")
    notes = wb.create_sheet("Notes")
    notes["A1"] = "Nothing here is computed."
    cached = {("Table", f"C{i}"): s / total for i, s in zip((3, 4, 5), (10, 30, 60), strict=True)}
    return inject_cached_values(workbook_bytes(wb), cached)


def test_empty_and_merged_sheets_flow_through_the_pipeline(client: TestClient) -> None:
    body = _upload(client, "odd.xlsx", _odd_workbook())
    assert body["status"] == "pending_review"
    assert body["validation_status"] in ("passed", "passed_with_warnings"), body["pipeline_notes"]
    vid = body["id"]
    model = client.get(f"/api/workbooks/{vid}/model").json()
    roles = {s["name"]: s["role"] for s in model["sheets"] if s["in_scope"]}
    assert set(roles) == {"Table", "Empty", "Notes"}
    run = client.post(f"/api/workbooks/{vid}/runs", json={}).json()
    values = client.get(f"/api/runs/{run['id']}/values", params={"sheet": "Table"}).json()["cells"]
    share = {c["address"]: c["value"] for c in values if c["formula"]}
    assert share == {"C3": 0.1, "C4": 0.3, "C5": 0.6}
    # Empty sheets answer with empty windows, never errors.
    grid = client.get(f"/api/runs/{run['id']}/grid", params={"sheet": "Empty"}).json()
    assert grid["rows"] == [] and grid["total_rows"] == 0
    assert (
        client.get(f"/api/runs/{run['id']}/values", params={"sheet": "Empty"}).json()["count"] == 0
    )
    # The merged title still reads as a label above the table.
    table = client.get(f"/api/runs/{run['id']}/grid", params={"sheet": "Table"}).json()
    assert [c["label"] for c in table["columns"]] == ["Region", "Sales", "Share"]


def test_missing_scoped_sheet_is_refused_explicitly_and_ignored_from_config(
    client: TestClient,
) -> None:
    vid = _upload(client, "logic.xlsx", logic_fixture_xlsx_bytes())["id"]
    refused = client.post(f"/api/workbooks/{vid}/interpret", json={"sheets": ["Inputs", "Nope"]})
    assert refused.status_code == 422 and "Nope" in refused.json()["detail"]
    # The config's sheetScope names another workbook's sheets: they are ignored, all sheets go in.
    ok = client.post(f"/api/workbooks/{vid}/interpret", json={"sheets": None})
    assert ok.status_code == 200 and set(ok.json()["scope"]) >= set(SCOPE)


def test_unrelated_workbook_never_touches_the_active_version(client: TestClient) -> None:
    first = _upload(client, "master.xlsx", logic_fixture_xlsx_bytes())
    assert (
        client.post(f"/api/workbooks/{first['id']}/activate", json={}).json()["status"] == "active"
    )
    stranger = _upload(client, "unrelated.xlsx", _odd_workbook())
    assert stranger["status"] == "pending_review"
    assert any(n.startswith("diff vs active") for n in stranger["pipeline_notes"])
    statuses = {v["id"]: v["status"] for v in client.get("/api/workbooks").json()}
    assert statuses[first["id"]] == "active" and statuses[stranger["id"]] == "pending_review"
    diff = client.get(f"/api/workbooks/{first['id']}/diff/{stranger['id']}").json()
    assert diff["summary"]["structural_changes"] > 0  # every sheet added / removed, nothing crashed


def _client() -> TestClient:
    c = TestClient(create_app())
    c.__enter__()
    return c


def test_two_users_run_what_ifs_concurrently(client: TestClient) -> None:
    vid = _upload(client, "logic.xlsx", logic_fixture_xlsx_bytes())["id"]
    client.post(f"/api/workbooks/{vid}/interpret", json={"sheets": SCOPE})
    assert client.post(f"/api/workbooks/{vid}/runs", json={}).status_code == 201
    settings = get_settings()
    assert settings.max_concurrent_runs == 1
    started = threading.Barrier(2)

    def what_if(units: int) -> list[int]:
        c = _client()
        started.wait()
        codes = []
        for _ in range(20):
            r = c.post(f"/api/workbooks/{vid}/runs", json={"overrides": {"Inputs!B3": units}})
            codes.append(r.status_code)
            if r.status_code == 201:
                assert r.json()["overrides"] == {"Inputs!B3": units}
                break
            assert r.status_code == 409 and r.json()["detail"]["busy"] is True
            assert r.headers.get("Retry-After") == str(settings.run_busy_retry_after_s)
            time.sleep(0.05)
        return codes

    with ThreadPoolExecutor(2) as pool:
        a, b = pool.map(what_if, (150, 200))
    assert a[-1] == 201 and b[-1] == 201
    # Both what-ifs are persisted with their own overrides; the limiter never lost one.
    runs = client.get(f"/api/workbooks/{vid}/runs").json()
    assert {r["overrides"].get("Inputs!B3") for r in runs if r["overrides"]} == {150, 200}
    assert len(run_store.state_cache._items) <= settings.state_cache_entries


def test_new_version_activated_mid_session_leaves_the_other_user_consistent(
    client: TestClient,
) -> None:
    first = _upload(client, "master.xlsx", logic_fixture_xlsx_bytes())
    client.post(f"/api/workbooks/{first['id']}/interpret", json={"sheets": SCOPE})
    client.post(f"/api/workbooks/{first['id']}/validate")
    client.post(f"/api/workbooks/{first['id']}/activate", json={})
    user_a, user_b = _client(), _client()
    run_a = user_a.post(
        f"/api/workbooks/{first['id']}/runs", json={"overrides": {"Inputs!B3": 150}}
    )
    assert run_a.status_code == 201

    second = _upload(user_b, "master-v2.xlsx", logic_fixture_xlsx_bytes(threshold=130))
    user_b.post(f"/api/workbooks/{second['id']}/interpret", json={"sheets": SCOPE})
    user_b.post(f"/api/workbooks/{second['id']}/validate")
    assert (
        user_b.post(f"/api/workbooks/{second['id']}/activate", json={}).json()["status"] == "active"
    )

    # User A's version is immutable: their what-if still reads, and further runs still work.
    outputs = user_a.get(
        f"/api/workbooks/{first['id']}/outputs", params={"run_id": run_a.json()["id"]}
    )
    assert outputs.status_code == 200 and outputs.json()["run_id"] == run_a.json()["id"]
    again = user_a.post(
        f"/api/workbooks/{first['id']}/runs", json={"overrides": {"Inputs!B3": 160}}
    )
    assert again.status_code == 201 and again.json()["version_id"] == first["id"]
    statuses = {v["id"]: v["status"] for v in user_a.get("/api/workbooks").json()}
    assert statuses[first["id"]] == "validated" and statuses[second["id"]] == "active"
    # And user B's new version has its own baseline and outputs.
    assert user_b.get(f"/api/workbooks/{second['id']}/outputs").status_code == 200


@pytest.mark.parametrize("bad", [b"not a workbook at all", b"PK\x03\x04garbage"])
def test_corrupt_uploads_are_rejected_cleanly(client: TestClient, bad: bytes) -> None:
    response = client.post("/api/workbooks", files={"file": ("broken.xlsx", bad, XLSX)})
    assert response.status_code == 422
    assert "could not be read" in response.json()["detail"]
    assert all(v["filename"] != "broken.xlsx" for v in client.get("/api/workbooks").json())
