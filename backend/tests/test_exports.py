"""Exports: view descriptors, csv / xlsx / pdf writers, the xlsx round trip, background jobs."""

import csv
import io
import time

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.engine.values import round15
from app.exports import report as report_mod
from app.exports.descriptor import ExportColumn, ExportView, Provenance
from app.parser import load_raw_workbook
from tests.fixtures.logic_workbook import ROWS, logic_fixture_xlsx_bytes, sales_for

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
SCOPE = ["Inputs", "Lookup", "Series", "Outputs", "Semantics"]


@pytest.fixture
def version_id(client: TestClient) -> str:
    response = client.post(
        "/api/workbooks", files={"file": ("logic.xlsx", logic_fixture_xlsx_bytes(), XLSX)}
    )
    assert response.status_code == 201, response.text
    vid = response.json()["id"]
    assert client.post(f"/api/workbooks/{vid}/interpret", json={"sheets": SCOPE}).status_code == 200
    return vid


@pytest.fixture
def baseline(client: TestClient, version_id: str) -> dict:
    run = client.post(f"/api/workbooks/{version_id}/runs", json={})
    assert run.status_code == 201, run.text
    return run.json()


def _rows(text: str) -> list[list[str]]:
    return list(csv.reader(io.StringIO(text.lstrip("﻿"))))


def test_csv_exports_exactly_the_grid_window(client: TestClient, baseline: dict) -> None:
    grid = client.get(
        f"/api/runs/{baseline['id']}/grid", params={"sheet": "Series", "r1": 2, "r2": 4}
    ).json()
    response = client.get(
        f"/api/runs/{baseline['id']}/export",
        params={"format": "csv", "sheet": "Series", "window": "A2:G4"},
    )
    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith("text/csv")
    assert "Series" in response.headers["content-disposition"]
    rows = _rows(response.text)
    assert rows[0] == ["Row"] + [c["label"] for c in grid["columns"]]
    assert len(rows) == 1 + len(grid["rows"])
    assert rows[1][:5] == ["2", "1", str(sales_for(1)), str(sales_for(1)), "Low"]
    assert rows[1][7] == "FALSE"  # Above units: 97 > 120 is FALSE
    assert (
        client.get(f"/api/runs/{baseline['id']}/export", params={"format": "csv"}).status_code
        == 422
    )
    assert (
        client.get(
            f"/api/runs/{baseline['id']}/export", params={"format": "csv", "sheet": "Nope"}
        ).status_code
        == 404
    )


def test_xlsx_round_trip_matches_the_stored_run(
    client: TestClient, version_id: str, baseline: dict, tmp_path
) -> None:
    response = client.get(
        f"/api/runs/{baseline['id']}/export",
        params={"format": "xlsx", "scope": "all", "formulas": "true"},
    )
    assert response.status_code == 200, response.text
    assert response.headers["content-type"] == XLSX
    wb = load_workbook(io.BytesIO(response.content), data_only=False)
    assert wb.sheetnames[0] == "Cover"
    cover = {
        row[0].value: row[1].value
        for row in wb["Cover"].iter_rows(min_row=3, max_row=12)
        if row[0].value
    }
    assert cover["Overrides"] == "none (baseline)"
    assert cover["Validation"] in ("passed", "passed with warnings", "failed")
    assert set(SCOPE) <= set(wb.sheetnames)
    assert "Series-formulas" in wb.sheetnames
    assert wb["Series-formulas"]["D3"].value.startswith("=IF(")

    # Re-ingest: every formula cell's value equals the run's stored value, in place.
    path = tmp_path / "export.xlsx"
    path.write_bytes(response.content)
    raw = load_raw_workbook(path)
    for sheet in SCOPE:
        stored = client.get(f"/api/runs/{baseline['id']}/values", params={"sheet": sheet}).json()[
            "cells"
        ]
        exported = raw.sheet(sheet)
        assert exported is not None
        for cell in stored:
            if not cell["formula"]:
                continue
            got = exported.cell(cell["address"])
            assert got is not None, (sheet, cell["address"])
            expected = cell["value"]
            if isinstance(expected, float):
                # openpyxl writes 16 significant digits; Excel compares at 15 (the platform's rule).
                assert round15(got.value) == round15(expected), (sheet, cell["address"])
            else:
                assert got.value == expected, (sheet, cell["address"], got.value, expected)
            if cell["type"] == "error":
                assert got.value_type == "error"


def test_xlsx_default_scope_is_the_output_sheets(client: TestClient, baseline: dict) -> None:
    response = client.get(f"/api/runs/{baseline['id']}/export", params={"format": "xlsx"})
    assert response.status_code == 200
    wb = load_workbook(io.BytesIO(response.content), read_only=True)
    # Output sheets by role: Outputs, and Semantics (nothing reads it), never the calculation sheets.
    assert wb.sheetnames == ["Cover", "Outputs", "Semantics"]
    ws = wb["Outputs"]
    values = {
        f"{c.column_letter}{c.row}": c.value
        for row in ws.iter_rows()
        for c in row
        if c.value is not None
    }
    assert values["A1"] == "Total sales" and values["B1"] == sum(
        sales_for(i) for i in range(1, ROWS + 1)
    )


def test_large_xlsx_exports_run_as_a_job(client: TestClient, baseline: dict) -> None:
    accepted = client.get(
        f"/api/runs/{baseline['id']}/export",
        params={"format": "xlsx", "scope": "all", "background": "true"},
    )
    assert accepted.status_code == 202, accepted.text
    job = accepted.json()
    assert job["status"] in ("running", "ok") and job["format"] == "xlsx"
    for _ in range(200):
        job = client.get(f"/api/exports/{job['id']}").json()
        if job["status"] != "running":
            break
        time.sleep(0.05)
    assert job["status"] == "ok", job
    assert job["filename"].endswith("-all-" + baseline["id"][:8] + ".xlsx")
    file = client.get(f"/api/exports/{job['id']}/file")
    assert file.status_code == 200 and file.headers["content-type"] == XLSX
    assert load_workbook(io.BytesIO(file.content), read_only=True).sheetnames[0] == "Cover"
    assert client.get("/api/exports/nope").status_code == 404


def test_changelog_exports_through_the_same_descriptor(client: TestClient) -> None:
    ids = []
    for name, variant in (("a.xlsx", {}), ("b.xlsx", {"threshold": 130})):
        r = client.post(
            "/api/workbooks", files={"file": (name, logic_fixture_xlsx_bytes(**variant), XLSX)}
        )
        vid = r.json()["id"]
        client.post(f"/api/workbooks/{vid}/interpret", json={"sheets": SCOPE})
        ids.append(vid)
    assert client.get(f"/api/workbooks/{ids[0]}/diff/{ids[1]}").status_code == 200
    csv_out = client.get(f"/api/workbooks/{ids[0]}/diff/{ids[1]}/export", params={"format": "csv"})
    assert csv_out.status_code == 200
    rows = _rows(csv_out.text)
    assert rows[0][:3] == ["#", "Group", "Change"]
    assert any(r[1] == "LOGIC" and "130" in r[6] for r in rows[1:]), rows[:3]
    xlsx_out = client.get(
        f"/api/workbooks/{ids[0]}/diff/{ids[1]}/export", params={"format": "xlsx"}
    )
    wb = load_workbook(io.BytesIO(xlsx_out.content), read_only=True)
    assert wb.sheetnames == ["Cover", "Version-changelog"]


def test_report_html_and_pdf(client: TestClient, version_id: str, baseline: dict) -> None:
    inc = client.post(
        f"/api/workbooks/{version_id}/runs", json={"overrides": {"Inputs!B3": 200}}
    ).json()
    html_text = client.get(f"/api/runs/{inc['id']}/report.html").text
    assert "Analysis report" in html_text and "what-if with 1 override(s)" in html_text
    assert "Total sales" in html_text and "Inputs!B3" in html_text
    assert "How these numbers are decided" in html_text
    assert "<table" in html_text
    pdf = client.get(f"/api/runs/{inc['id']}/export", params={"format": "pdf"})
    try:
        import weasyprint  # noqa: F401

        available = True
    except Exception:  # noqa: BLE001
        available = False
    if available:
        assert pdf.status_code == 200, pdf.text
        assert pdf.content.startswith(b"%PDF")
        assert pdf.headers["content-type"] == "application/pdf"
    else:
        assert pdf.status_code == 501
        assert "WeasyPrint" in pdf.json()["detail"]


def _view(rows, columns, labels=None, deltas=None) -> ExportView:
    prov = Provenance(
        filename="f.xlsx",
        version_id="v" * 32,
        version_status="active",
        exported_at=__import__("datetime").datetime.now(),
        scope="t",
    )
    return ExportView(
        id="t",
        title="t",
        columns=[ExportColumn(**c) for c in columns],
        row_labels=labels or [str(i) for i in range(len(rows))],
        rows=rows,
        deltas=deltas,
        provenance=prov,
    )


def test_report_summaries_are_category_level_and_top_n() -> None:
    columns = [
        {"key": "A", "label": "Name", "kind": "static"},
        {"key": "B", "label": "Category", "kind": "formula", "format": "general"},
        {"key": "C", "label": "Score", "kind": "output", "format": "number"},
    ]
    rows = [[f"fund {i}", "Alpha" if i % 3 else "Beta", float(i)] for i in range(1, 61)]
    view = _view(rows, columns, labels=[r[0] for r in rows])
    summary = report_mod.category_summary(view)
    assert summary["category"] == "Category" and summary["metric"] == "Score"
    assert [(k, n) for k, n, _m, _c in summary["rows"]] == [("Alpha", 40), ("Beta", 20)]
    # A finer classification beats a coarse flag; a name column (all distinct) never qualifies.
    finer = [{"key": "D", "label": "Flag", "kind": "formula"}] + columns
    rows2 = [
        ["Yes" if i % 2 else "No", r[0], ["Alpha", "Beta", "Gamma"][i % 3], r[2]]
        for i, r in enumerate(rows)
    ]
    assert (
        report_mod.category_summary(_view(rows2, finer, labels=[r[1] for r in rows2]))["category"]
        == "Category"
    )
    top = report_mod.top_rows(view, n=5)
    assert top["metric"] == "Score" and [r[0] for r in top["rows"]] == [
        f"fund {i}" for i in (60, 59, 58, 57, 56)
    ]
    deltas = [[None, None, (2.0 if i == 7 else -1.0 if i == 30 else None)] for i in range(1, 61)]
    moves = report_mod.largest_moves(
        _view(rows, columns, labels=[r[0] for r in rows], deltas=deltas)
    )
    assert moves["changed"] == 2 and moves["rows"][0][0] == "fund 7" and moves["rows"][0][4] == 2.0
    assert report_mod.category_summary(_view([], columns)) is None


def test_series_svg_draws_a_polyline() -> None:
    from app.storage.views import SeriesColumn, SeriesOut

    series = SeriesOut(
        x_label="Date", x_type="date", rows=3,
        columns=[SeriesColumn(label="Return", points=[["2025-01-01", 0.0], ["2025-01-02", 0.5], ["2025-01-03", 1.0]])],
    )  # fmt: skip
    svg = report_mod.series_svg(series)
    assert svg.startswith("<svg") and "<polyline" in svg and "2025-01-03" in svg
