from pathlib import Path

import pytest
from openpyxl.utils.datetime import to_excel

from app.parser import RawSheet, RawWorkbook, UnsupportedFileError, load_raw_workbook
from app.parser.inventory import FunctionInventory, functions_in_formula
from app.parser.loader import MACRO_WARNING
from tests.fixtures.workbook import (
    FIXTURE_CACHED_VALUES,
    FIXTURE_FORMULAS,
    START_DATE,
    as_xlsm,
    fixture_xlsx_bytes,
)


@pytest.fixture(scope="module")
def fixture_path(tmp_path_factory: pytest.TempPathFactory) -> Path:
    path = tmp_path_factory.mktemp("wb") / "fixture.xlsx"
    path.write_bytes(fixture_xlsx_bytes())
    return path


@pytest.fixture(scope="module")
def raw(fixture_path: Path) -> RawWorkbook:
    return load_raw_workbook(fixture_path)


def test_sheet_metadata(raw: RawWorkbook) -> None:
    assert raw.sheet_names == ["Inputs", "Lookup", "Calc"]
    assert [s.index for s in raw.sheets] == [0, 1, 2]
    assert raw.sheet("Lookup").state == "hidden"
    assert raw.sheet("Inputs").state == "visible"
    assert raw.sheet("Inputs").used_range == "A1:B6"
    assert raw.epoch == 1900
    assert raw.has_macros is False
    assert raw.warnings == []


def test_constants_and_inferred_types(raw: RawWorkbook) -> None:
    inputs = raw.sheet("Inputs")
    assert inputs.cell("A2").value == "Threshold"
    assert inputs.cell("A2").value_type == "text"
    assert inputs.cell("B2").value == 0.75
    assert inputs.cell("B2").value_type == "number"
    assert inputs.cell("B2").number_format == "0.00%"
    assert inputs.cell("B3").value == 120
    date_cell = inputs.cell("B4")
    assert date_cell.value_type == "date"
    assert date_cell.value == to_excel(START_DATE)
    assert date_cell.number_format == "yyyy-mm-dd"
    assert inputs.cell("B5").value is True
    assert inputs.cell("B5").value_type == "bool"
    assert inputs.cell("B6").value_type == "text"
    assert inputs.cell("Z99") is None
    for cell in inputs.records():
        assert cell.formula is None


def test_formulas_round_trip(raw: RawWorkbook) -> None:
    calc = raw.sheet("Calc")
    for addr, formula in FIXTURE_FORMULAS.items():
        assert calc.cell(addr).formula == formula, addr
    assert calc.formula_count == len(FIXTURE_FORMULAS)
    assert dict(calc.formulas()) == FIXTURE_FORMULAS


def test_cached_values_captured_for_every_formula_cell(raw: RawWorkbook) -> None:
    calc = raw.sheet("Calc")
    for (sheet, addr), expected in FIXTURE_CACHED_VALUES.items():
        cell = raw.sheet(sheet).cell(addr)
        assert cell.value == expected, addr
    assert calc.cell("B2").value_type == "number"
    assert calc.cell("B3").value_type == "text"
    assert calc.cell("B7").value_type == "error"
    assert calc.cell("B7").value == "#DIV/0!"
    assert calc.cell("B12").value_type == "bool"
    assert all(c.value_type != "empty" for c in calc.records() if c.formula)


def test_array_formula_expanded_to_every_cell(raw: RawWorkbook) -> None:
    calc = raw.sheet("Calc")
    for addr in ("B9", "B10"):
        cell = calc.cell(addr)
        assert cell.formula == "=Inputs!B2:B3*2"
        assert cell.array_ref == "B9:B10"
    assert calc.cell("B2").array_ref is None


def test_defined_names(raw: RawWorkbook) -> None:
    names = {n.name: n for n in raw.defined_names}
    assert names["Threshold"].refers_to == "Inputs!$B$2"
    assert names["Threshold"].scope is None
    assert names["Threshold"].builtin is False


def test_tables_and_merged_ranges(raw: RawWorkbook) -> None:
    assert [t.display_name for t in raw.tables] == ["Ratings"]
    table = raw.tables[0]
    assert table.sheet == "Lookup"
    assert table.ref == "A1:B4"
    assert table.columns == ["Code", "Rating"]
    lookup = raw.sheet("Lookup")
    assert lookup.tables == ["Ratings"]
    assert lookup.cell("A2").table == "Ratings"
    assert lookup.cell("B4").table == "Ratings"

    calc = raw.sheet("Calc")
    assert calc.merged_ranges == ["C1:D1"]
    assert calc.cell("C1").merged_range == "C1:D1"
    assert calc.cell("A1").merged_range is None


def test_function_inventory_is_exact(raw: RawWorkbook) -> None:
    # SUM appears inside a string literal in B8 and must not be counted twice.
    assert raw.functions == {
        "CONCATENATE": 1,
        "IF": 1,
        "IFERROR": 1,
        "SUM": 1,
        "VLOOKUP": 1,
        "XLOOKUP": 1,
    }


def test_summary_counts(raw: RawWorkbook) -> None:
    s = raw.summary()
    assert s.sheets == 3
    assert s.formula_cells == len(FIXTURE_FORMULAS)
    assert s.cells == sum(sheet.cell_count for sheet in raw.sheets)
    assert s.cells == 12 + 8 + (
        3 + 11 + 11
    )  # Inputs + Lookup + Calc(headers incl. merged, labels, formulas)
    assert s.named_ranges == 1
    assert s.tables == 1
    assert s.distinct_functions == 6
    assert list(s.functions)[0] in {"CONCATENATE", "IF", "IFERROR", "SUM", "VLOOKUP", "XLOOKUP"}
    assert s.sheet_names == ["Inputs", "Lookup", "Calc"]


def test_streaming_callback_and_keep_cells_false(fixture_path: Path) -> None:
    seen: list[RawSheet] = []
    raw = load_raw_workbook(fixture_path, on_sheet=seen.append, keep_cells=False)
    assert [s.name for s in seen] == ["Inputs", "Lookup", "Calc"]
    assert all(len(s.cells) == s.cell_count for s in seen)
    assert all(len(s.cells) == 0 for s in raw.sheets)
    assert [s.cell_count for s in raw.sheets] == [s.cell_count for s in seen]
    assert raw.summary().formula_cells == len(FIXTURE_FORMULAS)


def test_json_round_trip(raw: RawWorkbook) -> None:
    again = RawWorkbook.model_validate_json(raw.model_dump_json())
    assert again.model_dump() == raw.model_dump()
    assert again.sheet("Calc").cell("B7").value == "#DIV/0!"


def test_xlsm_is_ingested_and_flagged(tmp_path: Path) -> None:
    path = tmp_path / "macros.xlsm"
    path.write_bytes(as_xlsm(fixture_xlsx_bytes()))
    raw = load_raw_workbook(path)
    assert raw.has_macros is True
    assert MACRO_WARNING in raw.warnings
    assert raw.summary().formula_cells == len(FIXTURE_FORMULAS)


def test_xls_and_unknown_suffixes_are_rejected(tmp_path: Path) -> None:
    legacy = tmp_path / "old.xls"
    legacy.write_bytes(b"\xd0\xcf\x11\xe0")
    with pytest.raises(UnsupportedFileError, match="save as .xlsx"):
        load_raw_workbook(legacy)
    other = tmp_path / "notes.csv"
    other.write_text("a,b")
    with pytest.raises(UnsupportedFileError, match="Unsupported file type"):
        load_raw_workbook(other)


def test_formula_without_cached_values_reports_empty(tmp_path: Path) -> None:
    path = tmp_path / "nocache.xlsx"
    path.write_bytes(fixture_xlsx_bytes(with_cached_values=False))
    raw = load_raw_workbook(path)
    calc = raw.sheet("Calc")
    assert calc.cell("B2").formula == "=SUM(Inputs!B2:B3)"
    assert calc.cell("B2").value is None
    assert calc.cell("B2").value_type == "empty"


@pytest.mark.parametrize(
    ("formula", "expected"),
    [
        ('=IF(A1>0,"SUM(",_xlfn.XLOOKUP(1,B:B,C:C))', ["IF", "XLOOKUP"]),
        ("=Sheet1!A1+Table1[Col]", []),
        ('=SUM(A1:A3)+"x"&SUM(B1)', ["SUM", "SUM"]),
        ("=_xlfn._xlws.FILTER(A:A,B:B>1)", ["FILTER"]),
        ('=CONCATENATE("a ""quoted"" SUM(", B1)', ["CONCATENATE"]),
        ("=MyName*2", []),
        ("=sum(a1)", ["SUM"]),
    ],
)
def test_functions_in_formula(formula: str, expected: list[str]) -> None:
    assert functions_in_formula(formula) == expected


def test_inventory_counts_formulas_not_calls() -> None:
    inv = FunctionInventory()
    inv.add_formulas(["=SUM(A1)+SUM(B1)", "=IF(SUM(A1)>1,1,0)"])
    assert inv.as_dict() == {"SUM": 2, "IF": 1}
