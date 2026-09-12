"""Research views: the semantic map, the table, the endpoints, insights, movement, exports."""

import csv
import io
import re

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import dashboard_config
from app.config import get_settings
from app.research import table as table_store
from tests.fixtures.research_workbook import (
    CATEGORIES,
    expected,
    research_fixture_xlsx_bytes,
    research_map,
    write_config,
)

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@pytest.fixture(autouse=True)
def _fresh_database(tmp_path, monkeypatch):
    """Each research test starts from an empty store, so 'previous version' means what the
    test uploaded, not what an earlier test left behind."""
    from app.storage.db import get_engine, get_session_factory

    monkeypatch.setenv("MFA_DATA_DIR", str(tmp_path / "data"))
    for cached in (get_settings, get_engine, get_session_factory):
        cached.cache_clear()
    table_store.table_cache.clear()
    yield
    for cached in (get_settings, get_engine, get_session_factory):
        cached.cache_clear()


@pytest.fixture
def config(tmp_path, monkeypatch):
    """Point the app at a temp dashboard.config.json; yields a writer to swap the research map."""

    def use(research):
        path = write_config(tmp_path / "dashboard.config.json", research)
        monkeypatch.setenv("MFA_DASHBOARD_CONFIG_PATH", str(path))
        get_settings.cache_clear()
        dashboard_config._load.cache_clear()
        table_store.table_cache.clear()
        return path

    use(research_map())
    yield use
    monkeypatch.delenv("MFA_DASHBOARD_CONFIG_PATH", raising=False)
    get_settings.cache_clear()
    dashboard_config._load.cache_clear()
    table_store.table_cache.clear()


def _upload(client: TestClient, name: str, variant: str = "base", activate: bool = True) -> str:
    r = client.post(
        "/api/workbooks", files={"file": (name, research_fixture_xlsx_bytes(variant), XLSX)}
    )
    assert r.status_code == 201, r.text
    body = r.json()
    assert body["validation_status"] in ("passed", "passed_with_warnings"), body["pipeline_notes"]
    if activate:
        assert client.post(f"/api/workbooks/{body['id']}/activate", json={}).status_code == 200
    return body["id"]


def test_not_configured_paths(client: TestClient, config) -> None:
    config(None)
    _upload(client, "master.xlsx")
    body = client.get("/api/research/summary").json()
    assert body["configured"] is False and "absent" in body["problems"][0]
    assert client.get("/api/research/entities/export", params={"format": "csv"}).status_code == 409
    cfg = client.get("/api/research/config").json()
    assert cfg["configured"] is False

    bad = research_map()
    bad["measures"][7]["column"] = "ZZ"  # ret1y off the Perf sheet
    bad["measures"].append(
        {
            "key": "ghost",
            "label": "Ghost",
            "role": "return",
            "sheet": "Nope",
            "column": "B",
            "keyColumn": "A",
        }
    )
    config(bad)
    body = client.get("/api/research/summary").json()
    assert body["configured"] is True
    joined = " | ".join(body["problems"])
    assert "measures[7] 'ret1y': column ZZ is outside sheet 'Perf'" in joined
    assert "measures[10] 'ghost': sheet 'Nope' is not in scope" in joined
    cfg = client.get("/api/research/config").json()
    assert {m["key"]: m["resolved"] for m in cfg["measures"]}["ret1y"] is False

    wrong_entity = research_map()
    wrong_entity["entity"]["sheet"] = "Fund"
    config(wrong_entity)
    body = client.get("/api/research/entities").json()
    assert body["configured"] is False and "did you mean 'Funds'" in body["problems"][0]


def test_summary_counts_and_distribution(client: TestClient, config) -> None:
    _upload(client, "master.xlsx")
    body = client.get("/api/research/summary").json()
    assert body["configured"] is True and body["problems"] == []
    exp = expected()
    rated = sum(1 for e in exp.values() if e["quartile"] != "--")
    assert body["universe"]["total"] == 12
    assert body["universe"]["rated"] == rated == 8
    assert body["universe"]["categories"] == 3 and body["universe"]["unranked_categories"] == 1
    assert sum(body["quartiles"].values()) == rated
    assert body["category_averages"]["measure"] == "ret1y"
    assert body["validation"]["status"] in ("passed", "passed_with_warnings")
    assert body["findings"] == {"open": 0, "fixed": 1}
    assert body["movement"] is None and body["narrative"] is None  # one version: nothing to compare
    assert body["held_q1"]["available"] is False and "second genuine" in body["held_q1"]["note"]
    assert body["universe"]["unrated_small_categories"] == 3  # the three Direct-Beta funds
    assert body["universe"]["unrated_missing_data"] == 1  # Gamma One - Dir reads "--"
    assert body["universe_narrative"] == (
        "8 of 12 funds are rated. 1 of 3 categories has fewer than 4 funds (3 funds). "
        "1 fund is unrated for missing data: 0 too young (first NAV after 11 Feb 2016) and 1 with a data gap."
    )
    assert body["footer"] == "Test console · confidential"


def test_entities_filter_sort_page_and_pivot(client: TestClient, config) -> None:
    _upload(client, "master.xlsx")
    body = client.get("/api/research/entities").json()
    assert body["total"] == 12 and body["sort"] == "rank" and body["dir"] == "asc"
    ranks = [r["measures"]["rank"] for r in body["rows"]]
    assert ranks[:3] == [1, 1, 1] and ranks[-1] is None  # unranked last
    assert body["facets"]["category"] == sorted(CATEGORIES)
    assert body["rows"][0]["delta"] is None

    direct_q1 = client.get(
        "/api/research/entities", params={"plan": "Direct", "quartile": [1]}
    ).json()
    assert {r["dims"]["plan"] for r in direct_q1["rows"]} == {"Direct"}
    assert all(r["measures"]["quartile"] == 1 for r in direct_q1["rows"])
    search = client.get("/api/research/entities", params={"q": "gamma"}).json()
    assert {r["key"] for r in search["rows"]} == {"Gamma One - Dir", "Gamma One - Reg"}
    by_amc = client.get(
        "/api/research/entities", params={"amc": "Beta AMC", "sort": "expense", "dir": "asc"}
    ).json()
    assert [r["measures"]["expense"] for r in by_amc["rows"]] == sorted(
        r["measures"]["expense"] for r in by_amc["rows"]
    )
    page = client.get("/api/research/entities", params={"size": 5, "page": 3}).json()
    assert page["total"] == 12 and len(page["rows"]) == 2
    keys = client.get(
        "/api/research/entities", params={"keys": ["Alpha One - Dir", "Beta One - Reg"]}
    ).json()
    assert keys["total"] == 2

    pivot = client.get("/api/research/entities", params={"groupBy": "category"}).json()
    assert [g["key"] for g in pivot["groups"]] == sorted(CATEGORIES)
    assert sum(g["count"] for g in pivot["groups"]) == 12
    assert pivot["groups"][0]["subtotals"]["q1"] >= 1
    assert [r["group"] for r in pivot["rows"]][:5] == ["Direct-Alpha"] * 5
    bands = client.get("/api/research/entities", params={"groupBy": "corpus_band"}).json()
    assert len(bands["groups"]) == 4 and bands["groups"][0]["label"].startswith("Up to")
    assert client.get("/api/research/entities", params={"groupBy": "nope"}).status_code == 422


def test_entity_detail(client: TestClient, config) -> None:
    _upload(client, "master.xlsx")
    body = client.get("/api/research/entities/Beta One - Dir").json()
    assert body["configured"] is True
    m = {x["key"]: x for x in body["measures"]}
    assert m["rank"]["value"] == 1 and m["quartile"]["value"] == 1
    assert (
        m["score"]["categoryRank"] == 1 and m["score"]["categoryCount"] == 4
    )  # Gamma One has "--"
    assert m["ret1y"]["cell"] == "Perf!B4" and m["score"]["cell"] == "Funds!I4"
    assert [p["label"] for p in body["phases"]] == [
        "Bull 2020-21",
        "Bull 2022-24",
        "Bear 2020",
        "Bear 2022",
    ]
    assert body["phases"][0]["value"] == 22 and body["phases"][0]["categoryMean"] == pytest.approx(
        22.0
    )
    assert [p["label"] for p in body["periods"]] == ["1Y", "3Y"]
    assert body["peers"][0]["key"] == "Beta One - Dir" and body["peers"][0]["me"] is True
    assert [p["rank"] for p in body["peers"]] == sorted(
        p["rank"] for p in body["peers"] if p["rank"]
    ) + [None]
    assert body["category"]["count"] == 5 and body["category"]["rated"] == 4
    assert body["history"][0]["rank"] == 1 and len(body["history"]) == 1
    assert body["quartileExplanation"] and any("<=" in line for line in body["quartileExplanation"])
    assert body["narrative"] == "Beta One ranks 1 of 4 in Direct-Alpha (Q1)."  # no previous version
    assert client.get("/api/research/entities/Nobody").status_code == 404


def test_categories(client: TestClient, config) -> None:
    _upload(client, "master.xlsx")
    body = client.get("/api/research/categories", params={"measure": "ret3y"}).json()
    rows = {r["key"]: r for r in body["rows"]}
    assert set(rows) == set(CATEGORIES)
    assert rows["Direct-Beta"]["unranked"] is True and rows["Direct-Beta"]["rated"] == 0
    assert rows["Direct-Alpha"]["count"] == 5 and rows["Direct-Alpha"]["rated"] == 4
    assert sum(rows["Direct-Alpha"]["quartiles"].values()) == 4
    assert rows["Direct-Alpha"]["statLabels"] == {"B": "Average score"}
    assert rows["Direct-Alpha"]["stats"]["B"] == pytest.approx(
        sum(v for v in (76.0, 62.0, 88.0, 42.0)) / 4
    )
    assert body["statsOnly"] == 1  # Direct-Ghost is in the averages table with no fund behind it
    assert (
        body["narrative"].startswith(
            "3 categories, 1 unranked. Among the 2 categories with 2 or more rated funds, "
        )
        and "widest spread" in body["narrative"]
    )


def test_movement_labels_repairs(client: TestClient, config) -> None:
    defective = _upload(client, "master-old.xlsx", variant="repair")
    fixed = _upload(client, "master.xlsx")
    body = client.get("/api/research/movement").json()
    assert (
        body["available"] is True and body["from"]["id"] == defective and body["to"]["id"] == fixed
    )
    by_key = {m["key"]: m for m in body["risers"] + body["fallers"]}
    assert "Beta One - Dir" in by_key and by_key["Beta One - Dir"]["cause"] == "repair"
    assert body["repairs"][0]["key"] == "Beta One - Dir"
    assert "is a data repair" in body["narrative"] and "master" not in body["narrative"]
    assert body["from"]["date"] and body["to"]["date"]  # dates, never filenames

    shifted = _upload(client, "master-v3.xlsx", variant="shift")
    body = client.get("/api/research/movement").json()
    assert body["from"]["id"] == fixed and body["to"]["id"] == shifted
    risers = {m["key"]: m for m in body["risers"]}
    assert (
        risers["Alpha One - Dir"]["rankTo"] == 1 and risers["Alpha One - Dir"]["cause"] == "market"
    )
    assert any(m["key"] == "Beta One - Reg" for m in body["fallers"])
    explicit = client.get(
        "/api/research/movement", params={"from": defective, "to": shifted}
    ).json()
    assert explicit["from"]["id"] == defective and explicit["moved"] >= 2

    summary = client.get("/api/research/summary").json()
    assert summary["movement"]["previous"]["id"] == fixed and summary["movement"]["moved"] >= 2
    assert summary["held_q1"]["available"] is True and summary["held_q1"]["versions"] == 3
    assert "changed rank since" in summary["narrative"] and "master" not in summary["narrative"]
    assert "held Q1 across all 3 genuine uploads" in summary["narrative"]
    detail = client.get("/api/research/entities/Alpha One - Dir").json()
    assert detail["deltaLabel"].startswith("up") and detail["narrative"].startswith(
        "Alpha One ranks 1 of 4 in Direct-Alpha (Q1). Up 1 place since "
    )
    assert [h["rank"] for h in detail["history"]] == [1, 2, 1]  # repaired row, fixed, then shifted
    funds = client.get("/api/research/entities", params={"q": "Alpha One - Dir"}).json()
    assert funds["rows"][0]["delta"]["rank"] == 1


def test_insights_compute_sentences(client: TestClient, config) -> None:
    _upload(client, "master-old.xlsx", variant="repair")
    _upload(client, "master.xlsx")
    body = client.get("/api/research/insights").json()
    assert body["sections"] == ["winning", "houses", "cost", "coverage"]
    assert body["sectionLabels"] == {"winning": "Who is winning", "cost": "Cost and scale"}
    ins = {i["key"]: i for i in body["insights"]}
    assert all(i["problems"] == [] for i in ins.values()), [i["problems"] for i in ins.values()]
    assert ins["best_score"]["sentence"] == "11 funds carry a score; Beta One leads at 88.00."
    assert [r["label"] for r in ins["best_score"]["rows"]] == ["Beta One", "Beta One", "Alpha One"]
    assert ins["held_q1"]["sentence"].endswith("in all 2 genuine uploads.")
    assert ins["held_q1"]["count"] >= 1 and ins["held_q1"]["status"] == "ok"
    # Teams stay one group (attribute "team"); rows and sentences render the members readably.
    assert "B. Sen and E. Roy" in [r["label"] for r in ins["team_league"]["rows"]]
    assert ins["team_league"]["sentence"] in (
        "B. Sen manages 1 Q1 fund.",
        "The team of B. Sen and E. Roy manages 1 Q1 fund.",
    )
    # Coverage: the one "--" fund is old (a data gap), no fund is too young; predicates may test
    # dimensions and constants read from cells.
    assert ins["unrated_gap"]["sentence"] == "1 fund older than the earliest phase still shows --."
    assert [r["label"] for r in ins["unrated_gap"]["rows"]] == ["Gamma One"]
    assert ins["unrated_gap"]["rows"][0]["valueLabel"] == "27 Dec 2014"
    assert ins["unrated_young"]["count"] == 0 and ins["unrated_young"]["status"] == "ok"
    assert ins["unrated_young"]["sentence"] == "Every unrated fund is old enough to rate."
    assert ins["direct_leaders"]["count"] >= 1 and all(
        r["sub"].startswith("Direct") for r in ins["direct_leaders"]["rows"]
    )
    # A filter insight with a minimum category size counts only the eligible categories.
    assert ins["leaders"]["sentence"] == (
        "Across the 2 categories with 4 or more rated funds, Beta One leads Direct-Alpha."
    )
    # top with fraction 0.5 admits more funds than the default quarter.
    assert ins["cheap_half"]["count"] > ins["cheap_quarter"]["count"] >= 1
    assert ins["amc_league"]["sentence"].startswith("Beta AMC has the most Q1 funds")
    assert ins["best_value"]["count"] >= 1 and ins["best_value"]["rows"][0]["valueLabel"].endswith(
        "%"
    )
    assert re.match(
        r"^₹[\d,.]+ (lakh )?crore sits in \d+ funds ranked Q3 or Q4\.$",
        ins["corpus_at_risk"]["sentence"],
    )
    assert ins["dispersion"]["rows"][0]["valueLabel"].endswith("%")
    only = client.get("/api/research/insights", params={"section": "cost"}).json()
    assert {i["section"] for i in only["insights"]} == {"cost"}

    split = research_map()
    split["dimensions"]["manager"]["attribute"] = "person"
    config(split)
    ins = {i["key"]: i for i in client.get("/api/research/insights").json()["insights"]}
    labels = [r["label"] for r in ins["team_league"]["rows"]]
    assert "E. Roy" in labels and "B. Sen, E. Roy" not in labels

    broken = research_map()
    broken["insights"][0]["where"][0]["measure"] = "score_x"
    broken["insights"][2]["sentence"] = {
        "default": "{group.label} and {nonsense}",
        "zero": "{bogus}",
    }
    broken["insights"][1]["across"]["where"][0]["measure"] = "ret1y"
    broken["narratives"]["fund"] = [{"default": "{label}", "requires": ["nothing"]}]
    broken["insights"][4]["where"][0] = {"dimension": "nope", "op": "eq", "value": "x"}
    broken["insights"][5]["where"] = [{"measure": "score", "op": "gt", "value": "$ghost"}]
    broken["coverage"] = {"measure": "inception", "constant": "ghost"}
    config(broken)
    cfg = client.get("/api/research/config").json()
    joined = " | ".join(cfg["problems"])
    assert "measure 'score_x' is not declared; did you mean 'score'?" in joined
    assert "unknown sentence placeholder {nonsense}" in joined
    assert "unknown sentence placeholder {bogus}" in joined
    assert "across.where can only test the primary score, rank or quartile" in joined
    assert "narratives.fund: unknown placeholder {nothing}" in joined
    assert "dimension 'nope' is not declared" in joined
    assert "constant '$ghost' is not declared under constants" in joined
    assert "coverage: constant 'ghost' is not declared" in joined


def test_snapshots_make_same_month_uploads_one_genuine_version(client: TestClient, config) -> None:
    """Two activations of the same master are one genuine month: consistency cards wait for a
    second month, the movement narrative reads as a zero sentence, and both snapshots persist."""
    from sqlalchemy import select

    from app.storage.db import get_session_factory
    from app.storage.models import ResearchSnapshot

    first = _upload(client, "master.xlsx")
    second = _upload(client, "master-again.xlsx")
    with get_session_factory()() as session:
        rows = list(session.scalars(select(ResearchSnapshot)))
    assert {r.version_id for r in rows} == {first, second}
    assert len({r.fingerprint for r in rows}) == 1 and all(r.entity_count == 12 for r in rows)

    summary = client.get("/api/research/summary").json()
    assert summary["held_q1"]["available"] is False and summary["held_q1"]["versions"] == 1
    assert summary["movement"]["moved"] == 0 and summary["movement"]["same_month"] is True
    assert "data repair" not in summary["narrative"]  # requires moved > 0
    assert len(summary["history"]) == 1
    ins = {i["key"]: i for i in client.get("/api/research/insights").json()["insights"]}
    assert ins["held_q1"]["status"] == "unavailable" and ins["held_q1"]["sentence"] == ""
    assert "second genuine monthly upload" in ins["held_q1"]["note"]
    move = client.get("/api/research/movement").json()
    assert move["narrative"].startswith("No fund changed rank between the earlier upload of ")
    assert move["narrative"].endswith(" and the later upload of " + move["to"]["date"] + ".")
    assert summary["narrative"].startswith("No fund changed rank since the earlier upload of ")
    detail = client.get("/api/research/entities/Beta One - Dir").json()
    assert detail["narrative"].endswith("(Q1). Unchanged since " + move["from"]["date"] + ".")
    assert summary["universe"]["outside_universe"] == 0
    assert len(detail["history"]) == 1


def test_indian_currency_and_sentence_forms() -> None:
    from app.research.insights import fmt_inr_crore, render_sentence, render_sentences
    from app.research.semantic import SentenceSpec

    assert fmt_inr_crore(3911885) == "₹39.1 lakh crore"
    assert fmt_inr_crore(122954) == "₹1.2 lakh crore"
    assert fmt_inr_crore(86785) == "₹86,785 crore"
    assert fmt_inr_crore(45.26) == "₹45.3 crore"
    ctx = {"n": "1", "m": "12", "_n": {"n": 1, "m": 12}}
    assert (
        render_sentence("{n|fund|funds} and {m|fund|funds} {m?is|are}", ctx)
        == "1 fund and 12 funds are"
    )
    assert render_sentence("{missing} here", ctx) is None
    specs = [
        SentenceSpec(default="{z} moved.", zero="Nothing moved."),
        SentenceSpec(default="{z} repairs.", requires=["z"]),
        SentenceSpec(default="{n} fund held."),
    ]
    assert (
        render_sentences(specs, {"z": "0", "n": "1", "_n": {"z": 0, "n": 1}})
        == "Nothing moved. 1 fund held."
    )
    assert (
        render_sentences([SentenceSpec(default="{z} moved.")], {"z": "0", "_n": {"z": 0}}) is None
    )


def test_exports_go_through_the_view_descriptor(client: TestClient, config) -> None:
    _upload(client, "master.xlsx")
    csv_out = client.get(
        "/api/research/entities/export", params={"format": "csv", "plan": "Direct"}
    )
    assert csv_out.status_code == 200
    rows = list(csv.reader(io.StringIO(csv_out.text.lstrip("﻿"))))
    assert rows[0][:3] == ["Fund", "Identity", "Category"] and "Score" in rows[0]
    assert len(rows) == 1 + 8
    xlsx_out = client.get("/api/research/categories/export", params={"format": "xlsx"})
    wb = load_workbook(io.BytesIO(xlsx_out.content), read_only=True)
    assert wb.sheetnames == ["Cover", "Categories"]
    ins = client.get("/api/research/insights/export", params={"format": "csv"})
    assert ins.status_code == 200 and "Finding" in ins.text.splitlines()[0]
    assert client.get("/api/research/movement/export", params={"format": "csv"}).status_code == 409
    report = client.get(
        f"/api/runs/{client.get('/api/research/summary').json()['run_id']}/report.html"
    ).text
    assert "Insights brief" in report and "Beta One leads at 88.00" in report
