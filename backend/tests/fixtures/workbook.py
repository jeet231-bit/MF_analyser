"""Generated fixture workbook used across parser, model, engine and validation tests.

openpyxl never computes formulas, so the file it writes has no cached values. Excel does
write them, and the parser must capture them, so ``inject_cached_values`` patches the
sheet XML exactly the way Excel would have (``<v>`` elements with the right ``t`` type).
"""

from __future__ import annotations

import datetime as dt
import io
import re
import zipfile

import openpyxl
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.worksheet.table import Table

from app.parser.models import EXCEL_ERRORS
from app.parser.package import inspect_package

START_DATE = dt.date(2026, 1, 31)

# (sheet, address) -> cached value Excel would have stored.
FIXTURE_CACHED_VALUES: dict[tuple[str, str], object] = {
    ("Calc", "B2"): 120.75,
    ("Calc", "B3"): "High",
    ("Calc", "B4"): 2,
    ("Calc", "B5"): "n/a",
    ("Calc", "B6"): 1.5,
    ("Calc", "B7"): "#DIV/0!",
    ("Calc", "B8"): "SUM(Total)",
    ("Calc", "B9"): 1.5,
    ("Calc", "B10"): 240,
    ("Calc", "B11"): 1,
    ("Calc", "B12"): True,
}

FIXTURE_FORMULAS: dict[str, str] = {
    "B2": "=SUM(Inputs!B2:B3)",
    "B3": '=IF(Inputs!B3>100,"High","Low")',
    "B4": '=VLOOKUP("B",Lookup!A2:B4,2,FALSE)',
    "B5": '=IFERROR(1/0,"n/a")',
    "B6": "=Threshold*2",
    "B7": "=1/0",
    "B8": '=CONCATENATE("SUM(",A2,")")',
    "B9": "=Inputs!B2:B3*2",
    "B10": "=Inputs!B2:B3*2",
    "B11": '=_xlfn.XLOOKUP("A",Lookup!A2:A4,Lookup!B2:B4)',
    "B12": "=Inputs!B5",
}


def build_fixture_workbook() -> openpyxl.Workbook:
    wb = openpyxl.Workbook()

    inputs = wb.active
    inputs.title = "Inputs"
    inputs["A1"], inputs["B1"] = "Parameter", "Value"
    inputs["A2"], inputs["B2"] = "Threshold", 0.75
    inputs["A3"], inputs["B3"] = "Units", 120
    inputs["A4"], inputs["B4"] = "Start date", START_DATE
    inputs["B4"].number_format = "yyyy-mm-dd"
    inputs["A5"], inputs["B5"] = "Active", True
    inputs["A6"], inputs["B6"] = "Label", "Alpha"
    inputs["B2"].number_format = "0.00%"
    wb.defined_names["Threshold"] = DefinedName("Threshold", attr_text="Inputs!$B$2")

    lookup = wb.create_sheet("Lookup")
    lookup.sheet_state = "hidden"
    lookup["A1"], lookup["B1"] = "Code", "Rating"
    for row, (code, rating) in enumerate([("A", 1), ("B", 2), ("C", 3)], start=2):
        lookup[f"A{row}"], lookup[f"B{row}"] = code, rating
    lookup.add_table(Table(displayName="Ratings", ref="A1:B4"))

    calc = wb.create_sheet("Calc")
    calc["A1"], calc["B1"] = "Metric", "Result"
    calc.merge_cells("C1:D1")
    calc["C1"] = "Notes"
    labels = {
        2: "Total",
        3: "Flag",
        4: "Rating",
        5: "Safe",
        6: "Uses name",
        7: "Bad",
        8: "Text fn",
        9: "Array 1",
        10: "Array 2",
        11: "Future fn",
        12: "Bool",
    }
    for row, label in labels.items():
        calc[f"A{row}"] = label
    for addr, formula in FIXTURE_FORMULAS.items():
        if addr in ("B9", "B10"):
            continue
        calc[addr] = formula
    calc["B9"] = ArrayFormula("B9:B10", "=Inputs!B2:B3*2")
    calc["B2"].number_format = "#,##0.00"
    return wb


def workbook_bytes(wb: openpyxl.Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _xml_value(value: object) -> tuple[str, str]:
    """(t attribute or '', <v> text) mirroring Excel's serialisation."""
    if isinstance(value, bool):
        return ' t="b"', "1" if value else "0"
    if isinstance(value, int | float):
        return "", repr(value) if isinstance(value, float) else str(value)
    if isinstance(value, str) and value in EXCEL_ERRORS:
        return ' t="e"', value
    return ' t="str"', str(value)


def inject_cached_values(xlsx: bytes, values: dict[tuple[str, str], object]) -> bytes:
    """Return a copy of ``xlsx`` whose formula cells carry the given cached values."""
    info = inspect_package(io.BytesIO(xlsx))
    part_of = {s.name: s.path for s in info.sheets}
    by_part: dict[str, dict[str, object]] = {}
    for (sheet, addr), value in values.items():
        by_part.setdefault(part_of[sheet], {})[addr] = value

    src = zipfile.ZipFile(io.BytesIO(xlsx))
    out_buf = io.BytesIO()
    with zipfile.ZipFile(out_buf, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename in by_part:
                data = _patch_sheet_xml(data.decode("utf-8"), by_part[item.filename]).encode(
                    "utf-8"
                )
            dst.writestr(item, data)
    return out_buf.getvalue()


def _patch_sheet_xml(xml: str, values: dict[str, object]) -> str:
    for addr, value in values.items():
        t_attr, text = _xml_value(value)
        pattern = re.compile(
            rf'<c r="{addr}"((?: [^>]*)?)>(<f\b[^>]*>.*?</f>|<f\b[^>]*/>)<v ?/>', re.DOTALL
        )
        m = pattern.search(xml)
        if m is not None:
            attrs = re.sub(r' t="[^"]*"', "", m.group(1))
            xml = (
                xml[: m.start()]
                + f'<c r="{addr}"{attrs}{t_attr}>{m.group(2)}<v>{text}</v>'
                + xml[m.end() :]
            )
            continue
        # Non-anchor array cells have no <c> element yet; Excel writes a value-only cell.
        row = int(re.sub(r"[A-Z]+", "", addr))
        row_re = re.compile(rf'(<row r="{row}"[^>]*>.*?)(</row>)', re.DOTALL)
        rm = row_re.search(xml)
        if rm is None:
            raise AssertionError(
                f"row {row} not present in sheet XML; add a label cell so the row exists"
            )
        xml = (
            xml[: rm.start()]
            + rm.group(1)
            + f'<c r="{addr}"{t_attr}><v>{text}</v></c>'
            + rm.group(2)
            + xml[rm.end() :]
        )
    return xml


def fixture_xlsx_bytes(*, with_cached_values: bool = True) -> bytes:
    data = workbook_bytes(build_fixture_workbook())
    return inject_cached_values(data, FIXTURE_CACHED_VALUES) if with_cached_values else data


def build_logic_fixture_workbook_bytes() -> bytes:
    """Raw bytes of the extended (logic) fixture without cached values, for perturbation tests."""
    from tests.fixtures.logic_workbook import build_logic_fixture_workbook

    return workbook_bytes(build_logic_fixture_workbook())


def as_xlsm(xlsx: bytes) -> bytes:
    """Repackage as a macro-enabled workbook by adding a (dummy) VBA project part."""
    src = zipfile.ZipFile(io.BytesIO(xlsx))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            dst.writestr(item, src.read(item.filename))
        dst.writestr("xl/vbaProject.bin", b"\xd0\xcf\x11\xe0 dummy vba project")
    return out.getvalue()
