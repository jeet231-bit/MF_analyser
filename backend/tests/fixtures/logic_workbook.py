"""Extended fixture for the logic model: every classification and rule pattern in one file.

Builds on the Phase 1 fixture (Inputs / Lookup / Calc) and adds:
- Series: a 20-row table with copied formulas (one template per column), a running total
  (self-dependent block), a nested-IF band ladder, a fixed-table VLOOKUP, an IFERROR fallback,
  and a threshold against an absolute labelled cell on another sheet.
- Outputs: three formulas reading Series that nothing else reads (output role, block sinks).
- Calc gets a static note nothing references.
"""

from __future__ import annotations

from tests.fixtures.workbook import (
    FIXTURE_CACHED_VALUES,
    build_fixture_workbook,
    inject_cached_values,
    workbook_bytes,
)

ROWS = 20
SALES = [90 + 7 * i for i in range(1, ROWS + 1)]  # 97 .. 230
BANDS = {"Low": 1, "Mid": 2, "High": 3}
UNITS = 120  # Inputs!B3


def band(sales: int) -> str:
    return "Low" if sales <= 120 else ("Mid" if sales <= 200 else "High")


def build_logic_fixture_workbook():
    wb = build_fixture_workbook()
    lookup = wb["Lookup"]
    for i, (name, code) in enumerate(BANDS.items(), start=1):
        lookup[f"D{i}"], lookup[f"E{i}"] = name, code

    wb["Calc"]["A14"] = "Notes: values in INR lakh"

    s = wb.create_sheet("Series")
    for col, header in zip(
        "ABCDEFG",
        ["Month", "Sales", "Cumulative", "Band", "Rated", "Safe", "Above units"],
        strict=True,
    ):
        s[f"{col}1"] = header
    for i in range(1, ROWS + 1):
        r = i + 1
        s[f"A{r}"] = i
        s[f"B{r}"] = SALES[i - 1]
        s[f"C{r}"] = "=B2" if r == 2 else f"=C{r - 1}+B{r}"
        s[f"D{r}"] = f'=IF(B{r}<=120,"Low",IF(B{r}<=200,"Mid","High"))'
        s[f"E{r}"] = f"=VLOOKUP(D{r},Lookup!$D$1:$E$3,2,FALSE)"
        s[f"F{r}"] = f'=IFERROR(E{r}/B{r},"n/a")'
        s[f"G{r}"] = f"=B{r}>Inputs!$B$3"

    o = wb.create_sheet("Outputs")
    o["A1"], o["B1"] = "Total sales", "=SUM(Series!B2:B21)"
    o["A2"], o["B2"] = "Peak cumulative", "=MAX(Series!C2:C21)"
    o["A3"], o["B3"] = "High months", '=COUNTIF(Series!D2:D21,"High")'
    return wb


def logic_cached_values() -> dict[tuple[str, str], object]:
    values: dict[tuple[str, str], object] = dict(FIXTURE_CACHED_VALUES)
    cumulative = 0
    for i in range(1, ROWS + 1):
        r = i + 1
        sales = SALES[i - 1]
        cumulative += sales
        b = band(sales)
        rated = BANDS[b]
        values[("Series", f"C{r}")] = cumulative
        values[("Series", f"D{r}")] = b
        values[("Series", f"E{r}")] = rated
        values[("Series", f"F{r}")] = rated / sales
        values[("Series", f"G{r}")] = sales > UNITS
    values[("Outputs", "B1")] = sum(SALES)
    values[("Outputs", "B2")] = cumulative
    values[("Outputs", "B3")] = sum(1 for x in SALES if band(x) == "High")
    return values


def logic_fixture_xlsx_bytes() -> bytes:
    return inject_cached_values(
        workbook_bytes(build_logic_fixture_workbook()), logic_cached_values()
    )


def cycle_fixture_xlsx_bytes() -> bytes:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Cyc"
    ws["A1"], ws["B1"], ws["C1"] = "=B1+1", "=A1+1", 5
    ws["A2"] = "=C1*2"
    return inject_cached_values(
        workbook_bytes(wb), {("Cyc", "A1"): 0, ("Cyc", "B1"): 0, ("Cyc", "A2"): 10}
    )
