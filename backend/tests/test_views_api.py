"""Dashboard view endpoints: grid window, inputs catalogue, outputs summary, explanations,
background runs."""

import datetime as dt
import time

import openpyxl
import pytest
from fastapi.testclient import TestClient

from tests.fixtures.logic_workbook import ROWS, UNITS, logic_fixture_xlsx_bytes, sales_for
from tests.fixtures.workbook import inject_cached_values, workbook_bytes

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
SCOPE = ["Inputs", "Lookup", "Series", "Outputs", "Semantics"]


@pytest.fixture
def version_id(client: TestClient) -> str:
    response = client.post(
        "/api/workbooks", files={"file": ("logic.xlsx", logic_fixture_xlsx_bytes(), XLSX)}
    )
    assert response.status_code == 201, response.text
    vid = response.json()["id"]
    assert client.post(f"/api/workbooks/{vid}/interpret", json={"sheets": SCOPE}).status_code == 200
    return vid


@pytest.fixture
def baseline(client: TestClient, version_id: str) -> dict:
    run = client.post(f"/api/workbooks/{version_id}/runs", json={})
    assert run.status_code == 201, run.text
    return run.json()


def test_grid_window_mirrors_the_sheet_with_labels_and_markers(
    client: TestClient, version_id: str, baseline: dict
) -> None:
    grid = client.get(f"/api/runs/{baseline['id']}/grid", params={"sheet": "Series"}).json()
    assert grid["body_start"] == 2 and grid["r1"] == 2 and grid["r2"] == ROWS + 1
    assert [c["letter"] for c in grid["columns"]][:4] == ["A", "B", "C", "D"]
    labels = {c["letter"]: c["label"] for c in grid["columns"]}
    assert labels["B"] == "Sales" and labels["D"] == "Band"
    kinds = {c["letter"]: c["kind"] for c in grid["columns"]}
    assert kinds["A"] == "static" and kinds["B"] == "input" and kinds["C"] == "formula"
    first = grid["rows"][0]
    assert first["row"] == 2
    a, b, c, d = first["cells"][:4]
    assert a == {"v": 1, "t": "number"} and b["v"] == sales_for(1)
    assert c["f"] is True and c["v"] == sales_for(1)
    assert d["v"] == "Low"
    assert grid["baseline_run_id"] is None

    # Explicit window and caps.
    part = client.get(
        f"/api/runs/{baseline['id']}/grid",
        params={"sheet": "Series", "r1": 1, "r2": 3, "c1": 2, "c2": 3},
    ).json()
    assert (part["r1"], part["r2"], part["c1"], part["c2"]) == (1, 3, 2, 3)
    assert part["rows"][0]["cells"][0] == {"v": "Sales", "t": "text"}
    assert (
        client.get(f"/api/runs/{baseline['id']}/grid", params={"sheet": "Nope"}).status_code == 404
    )
    assert client.get("/api/runs/nope/grid", params={"sheet": "Series"}).status_code == 404


def test_grid_marks_overrides_and_baseline_deltas(
    client: TestClient, version_id: str, baseline: dict
) -> None:
    inc = client.post(
        f"/api/workbooks/{version_id}/runs", json={"overrides": {"Inputs!B3": 200}}
    ).json()
    assert inc["kind"] == "incremental"
    inputs = client.get(f"/api/runs/{inc['id']}/grid", params={"sheet": "Inputs", "r1": 1}).json()
    b3 = inputs["rows"][2]["cells"][1]
    assert b3 == {"v": 200, "t": "number", "o": True}
    series = client.get(f"/api/runs/{inc['id']}/grid", params={"sheet": "Series"}).json()
    assert series["baseline_run_id"] == baseline["id"]
    g_col = next(i for i, c in enumerate(series["columns"]) if c["letter"] == "G")
    changed = [r for r in series["rows"] if r["cells"][g_col] and "b" in r["cells"][g_col]]
    # Rows whose sales fall between 120 and 200 flip "Above units" from TRUE to FALSE.
    assert len(changed) == sum(1 for i in range(1, ROWS + 1) if UNITS < sales_for(i) <= 200)
    assert changed[0]["cells"][g_col]["b"] is True and changed[0]["cells"][g_col]["v"] is False


def test_inputs_catalogue_groups_blocks_with_labels(
    client: TestClient, version_id: str, baseline: dict
) -> None:
    body = client.get(f"/api/workbooks/{version_id}/inputs").json()
    by_sheet = {s["sheet"]: s for s in body["sheets"]}
    assert set(by_sheet) >= {"Inputs", "Lookup", "Series"}
    assert by_sheet["Inputs"]["role"] == "input"
    params = [b for b in by_sheet["Inputs"]["blocks"] if b["shape"] == "parameters"]
    assert params and all(b["editable"] for b in params)
    cells = {c["address"]: c for b in params for c in b["cells"]}
    assert cells["B2"]["label"] == "Threshold" and cells["B2"]["value"] == 0.75
    assert cells["B3"]["label"] == "Units" and cells["B3"]["value"] == UNITS
    assert cells["B3"]["override"] is False
    table = next(b for b in by_sheet["Series"]["blocks"] if b["shape"] == "table")
    assert table["cell_count"] == ROWS and table["column_labels"] == ["Sales"]
    assert table["cells"] == []  # tables page through /grid

    inc = client.post(
        f"/api/workbooks/{version_id}/runs", json={"overrides": {"Inputs!B3": 200}}
    ).json()
    with_run = client.get(
        f"/api/workbooks/{version_id}/inputs", params={"run_id": inc["id"]}
    ).json()
    cells = {
        c["address"]: c
        for s in with_run["sheets"]
        if s["sheet"] == "Inputs"
        for b in s["blocks"]
        for c in b["cells"]
    }
    assert cells["B3"]["value"] == 200 and cells["B3"]["override"] is True
    assert client.get("/api/workbooks/nope/inputs").status_code == 404


def test_outputs_summary_has_metrics_rules_and_deltas(
    client: TestClient, version_id: str, baseline: dict
) -> None:
    body = client.get(f"/api/workbooks/{version_id}/outputs").json()
    assert body["run_id"] == baseline["id"] and body["baseline_run_id"] is None
    names = [s["sheet"] for s in body["sheets"]]
    assert names[0] == "Outputs"
    outputs = body["sheets"][0]
    metrics = {m["label"]: m for m in outputs["metrics"]}
    assert metrics["Total sales"]["value"] == sum(sales_for(i) for i in range(1, ROWS + 1))
    assert metrics["High months"]["delta"] is None
    assert outputs["series"] is None
    assert any(b["kind"] == "metric" for b in outputs["blocks"])

    inc = client.post(
        f"/api/workbooks/{version_id}/runs", json={"overrides": {"Inputs!B3": 200}}
    ).json()
    what_if = client.get(
        f"/api/workbooks/{version_id}/outputs", params={"run_id": inc["id"]}
    ).json()
    assert what_if["baseline_run_id"] == baseline["id"]
    outputs = what_if["sheets"][0]
    metrics = {m["label"]: m for m in outputs["metrics"]}
    # Units only feeds Series!G (Above units), which no output metric reads: values unchanged.
    assert metrics["Total sales"]["baseline"] == metrics["Total sales"]["value"]
    assert metrics["Total sales"]["delta"] is None
    series_sheet = next((s for s in what_if["sheets"] if s["sheet"] == "Series"), None)
    assert series_sheet is None  # Series is a calculation sheet, not an output
    assert client.get("/api/workbooks/nope/outputs").status_code == 404


def _daily_workbook() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Daily"
    ws["A1"], ws["B1"], ws["C1"] = "Date", "Level", "Return"
    cached = {}
    start = dt.date(2025, 1, 1)
    for i in range(60):
        r = i + 2
        ws[f"A{r}"] = start + dt.timedelta(days=i)
        ws[f"B{r}"] = 100 + i
        ws[f"C{r}"] = f"=B{r}/B$2-1"
        cached[("Daily", f"C{r}")] = (100 + i) / 100 - 1
    return inject_cached_values(workbook_bytes(wb), cached)


def test_outputs_detect_a_time_series(client: TestClient) -> None:
    response = client.post(
        "/api/workbooks", files={"file": ("daily.xlsx", _daily_workbook(), XLSX)}
    )
    vid = response.json()["id"]
    assert client.post(f"/api/workbooks/{vid}/interpret", json={"sheets": None}).status_code == 200
    assert client.post(f"/api/workbooks/{vid}/runs", json={}).status_code == 201
    body = client.get(f"/api/workbooks/{vid}/outputs").json()
    daily = next(s for s in body["sheets"] if s["sheet"] == "Daily")
    series = daily["series"]
    assert series is not None and series["x_type"] == "date" and series["rows"] == 60
    assert series["x_label"] == "Date"
    assert [c["label"] for c in series["columns"]] == ["Return"]
    points = series["columns"][0]["points"]
    assert points[0] == ["2025-01-01", 0.0] and len(points) == 60
    assert daily["table_rect"] == "C2:C61"


def test_lineage_explains_if_ladders_and_fallbacks(
    client: TestClient, version_id: str, baseline: dict
) -> None:
    node = client.get(f"/api/workbooks/{version_id}/lineage/Series!D6").json()
    # Row 6 is month 5: sales 125 -> not <= 120, but <= 200 -> "Mid".
    assert node["value"] == "Mid"
    assert node["explanation"] == [
        "B6<=120: 125 <= 120 is FALSE",
        "B6<=200: 125 <= 200 is TRUE",
        'Result: "Mid"',
    ]
    assert node["rule"] and "Low" in node["rule"]
    safe = client.get(f"/api/workbooks/{version_id}/lineage/Series!F6").json()
    assert safe["explanation"] == ["E6/B6 did not error", "Result: 0.016"]
    missing = client.get(f"/api/workbooks/{version_id}/lineage/Semantics!C21").json()
    assert missing["explanation"] == ['C20 gave #N/A; fallback "missing" used', 'Result: "missing"']
    plain = client.get(f"/api/workbooks/{version_id}/lineage/Series!C6").json()
    assert "explanation" not in plain  # nothing to narrate for =C5+B6
    off = client.get(
        f"/api/workbooks/{version_id}/lineage/Series!D6", params={"explanation": "false"}
    ).json()
    assert "explanation" not in off and off["value"] == "Mid"


def test_lineage_explanation_follows_a_what_if_run(
    client: TestClient, version_id: str, baseline: dict
) -> None:
    inc = client.post(
        f"/api/workbooks/{version_id}/runs", json={"overrides": {"Inputs!B3": 200}}
    ).json()
    node = client.get(
        f"/api/workbooks/{version_id}/lineage/Series!G6", params={"run_id": inc["id"]}
    ).json()
    assert node["value"] is False  # 125 > 200 is FALSE under the override
    reads = {r["range"]: r for r in node["reads"]}
    assert reads["Inputs!B3"]["cells"][0]["value"] == 200 if "Inputs!B3" in reads else True


def test_background_run_is_polled_to_completion(client: TestClient, version_id: str) -> None:
    accepted = client.post(f"/api/workbooks/{version_id}/runs", json={"background": True})
    assert accepted.status_code == 202, accepted.text
    run = accepted.json()
    assert run["status"] in ("running", "ok")
    for _ in range(200):
        run = client.get(f"/api/runs/{run['id']}").json()
        if run["status"] != "running":
            break
        time.sleep(0.05)
    assert run["status"] == "ok", run
    assert run["error"] is None and run["summary"]["blocks_evaluated"] > 10
    values = client.get(f"/api/runs/{run['id']}/values", params={"sheet": "Outputs"}).json()
    assert values["count"] == 6
    # A completed background full run without overrides is a baseline like any other.
    assert client.get(f"/api/workbooks/{version_id}/outputs").json()["run_id"] == run["id"]


def test_background_run_failure_is_reported_on_the_row(client: TestClient, version_id: str) -> None:
    client.post(f"/api/workbooks/{version_id}/interpret", json={"sheets": None})  # includes XLOOKUP
    accepted = client.post(f"/api/workbooks/{version_id}/runs", json={"background": True})
    assert accepted.status_code == 202
    run = accepted.json()
    for _ in range(200):
        run = client.get(f"/api/runs/{run['id']}").json()
        if run["status"] != "running":
            break
        time.sleep(0.05)
    assert run["status"] == "failed"
    assert "XLOOKUP" in run["error"]
    statuses = [
        (r["kind"], r["status"]) for r in client.get(f"/api/workbooks/{version_id}/runs").json()
    ]
    # The upload pipeline's validation run (unsupported templates skipped) is still the baseline.
    assert statuses == [("full", "failed"), ("full", "ok")], statuses
    assert client.get(f"/api/workbooks/{version_id}/outputs").json()["run_id"] != run["id"]
