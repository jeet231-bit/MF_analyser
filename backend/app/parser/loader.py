"""Load an .xlsx/.xlsm into a RawWorkbook.

Two streaming passes with openpyxl (read_only=True): one with formulas, one with cached
values (data_only=True). Read-only mode is mandatory: real workbooks here exceed a million
non-empty cells and a full load takes minutes and gigabytes. What read-only mode omits
(merged ranges, tables, defined names, macros) comes from ``package.inspect_package``.
"""

from __future__ import annotations

import logging
import time
from collections.abc import Callable
from pathlib import Path

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter, range_boundaries
from openpyxl.utils.cell import coordinate_from_string
from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula

from app.parser.cellvalues import coerce_value
from app.parser.inventory import FunctionInventory
from app.parser.models import CellColumns, PivotSpec, RawSheet, RawWorkbook, SheetState
from app.parser.package import SheetPart, inspect_package
from app.parser.pivots import parse_pivots

log = logging.getLogger(__name__)

SUPPORTED_SUFFIXES = {".xlsx", ".xlsm"}
MACRO_WARNING = "contains macros — macro logic is not interpreted"
PIVOT_WARNING = "contains pivot tables — pivot outputs are read as cached values, not recalculated"
EXTERNAL_LINK_WARNING = (
    "contains external links — external references are treated as declared inputs"
)

SheetCallback = Callable[[RawSheet], None]


class UnsupportedFileError(ValueError):
    pass


def _cells_in_range(ref: str) -> list[str]:
    min_col, min_row, max_col, max_row = range_boundaries(ref)
    return [
        f"{get_column_letter(c)}{r}"
        for r in range(min_row, max_row + 1)
        for c in range(min_col, max_col + 1)
    ]


def _membership(ranges: dict[str, list[str]]) -> dict[str, str]:
    """address -> range/table name, for merged ranges and table refs."""
    out: dict[str, str] = {}
    for name, refs in ranges.items():
        for ref in refs:
            for addr in _cells_in_range(ref):
                out[addr] = name
    return out


def _formula_text(value: object) -> tuple[str | None, str | None]:
    """(formula, array_ref) for a formula-mode cell value."""
    if isinstance(value, ArrayFormula):
        text = value.text or ""
        return (text if text.startswith("=") else f"={text}"), value.ref
    if isinstance(value, DataTableFormula):
        return f"={{=TABLE({value.r1 or ''},{value.r2 or ''})}}", value.ref
    if isinstance(value, str) and value.startswith("="):
        return value, None
    return None, None


def _state(raw: str) -> SheetState:
    return raw if raw in ("visible", "hidden", "veryHidden") else "visible"  # type: ignore[return-value]


def _extract_sheet(
    ws_formulas,
    ws_values,
    part: SheetPart,
    index: int,
    epoch: int,
    inventory: FunctionInventory,
) -> RawSheet:
    cells = CellColumns()
    positions: dict[str, int] = {}
    merged_of = _membership({r: [r] for r in part.merged_ranges})
    table_of = _membership({t.display_name: [t.ref] for t in part.tables})
    array_anchors: list[tuple[str, str, str]] = []  # (anchor, formula, ref)
    formula_count = 0

    # Pass 1: formulas and constants.
    for row in ws_formulas.iter_rows():
        for c in row:
            raw = c.value
            if raw is None:
                continue
            addr = c.coordinate
            formula, array_ref = _formula_text(raw)
            if formula is not None:
                value, value_type = None, "empty"
                formula_count += 1
                inventory.add_formula(formula)
                if array_ref is not None:
                    array_anchors.append((addr, formula, array_ref))
            else:
                value, value_type = coerce_value(raw, c.data_type, epoch)
            positions[addr] = cells.append(
                addr,
                c.row,
                c.column,
                formula=formula,
                value=value,
                value_type=value_type,
                number_format=c.number_format or "General",
                merged_range=merged_of.get(addr),
                table=table_of.get(addr),
                array_ref=array_ref,
            )

    # Array formulas: every cell in the ref owns the anchor's formula.
    for anchor, formula, ref in array_anchors:
        for addr in _cells_in_range(ref):
            if addr == anchor:
                continue
            pos = positions.get(addr)
            if pos is None:
                col_letters, row_num = coordinate_from_string(addr)
                pos = cells.append(
                    addr,
                    row_num,
                    column_index_from_string(col_letters),
                    formula=formula,
                    value=None,
                    value_type="empty",
                    number_format="General",
                    merged_range=merged_of.get(addr),
                    table=table_of.get(addr),
                    array_ref=ref,
                )
                positions[addr] = pos
            else:
                cells.formula[pos] = formula
                cells.array_ref[pos] = ref
            formula_count += 1
            inventory.add_formula(formula)

    # Pass 2: cached values for formula cells.
    for row in ws_values.iter_rows():
        for c in row:
            raw = c.value
            if raw is None:
                continue
            pos = positions.get(c.coordinate)
            if pos is None or cells.formula[pos] is None:
                continue
            value, value_type = coerce_value(raw, c.data_type, epoch)
            cells.value[pos] = value
            cells.value_type[pos] = value_type

    used_range = None
    if cells.row:
        used_range = (
            f"{get_column_letter(min(cells.col))}{min(cells.row)}:"
            f"{get_column_letter(max(cells.col))}{max(cells.row)}"
        )

    return RawSheet(
        name=part.name,
        index=index,
        state=_state(part.state),
        used_range=used_range,
        merged_ranges=list(part.merged_ranges),
        tables=[t.display_name for t in part.tables],
        cell_count=len(cells),
        formula_count=formula_count,
        cells=cells,
    )


def load_raw_workbook(
    path: str | Path,
    *,
    filename: str | None = None,
    on_sheet: SheetCallback | None = None,
    keep_cells: bool = True,
) -> RawWorkbook:
    """Parse the file at ``path``.

    ``on_sheet`` is called with each fully extracted sheet; with ``keep_cells=False`` the
    returned workbook keeps only sheet metadata, which lets a caller persist sheets one at a
    time and bound peak memory to the largest sheet.
    """
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".xls":
        raise UnsupportedFileError(
            "Legacy binary .xls files are not supported. Open the file in Excel, save as .xlsx and upload again."
        )
    if suffix not in SUPPORTED_SUFFIXES:
        raise UnsupportedFileError(
            f"Unsupported file type '{suffix}'. Upload an .xlsx or .xlsm workbook."
        )

    started = time.perf_counter()
    info = inspect_package(path)
    warnings: list[str] = []
    if info.has_macros or suffix == ".xlsm":
        warnings.append(MACRO_WARNING)
    pivots: list[PivotSpec] = []
    if info.has_pivots:
        warnings.append(PIVOT_WARNING)
        pivots = parse_pivots(path)
    if info.has_external_links:
        warnings.append(EXTERNAL_LINK_WARNING)

    wb_f = openpyxl.load_workbook(path, read_only=True, data_only=False, keep_links=True)
    wb_v = openpyxl.load_workbook(path, read_only=True, data_only=True, keep_links=True)
    inventory = FunctionInventory()
    sheets: list[RawSheet] = []
    try:
        for index, part in enumerate(info.sheets):
            if part.is_chartsheet:
                warnings.append(f"chart sheet '{part.name}' skipped — chart sheets hold no cells")
                continue
            if part.name not in wb_f.sheetnames:
                warnings.append(
                    f"sheet '{part.name}' declared in workbook.xml but not readable — skipped"
                )
                continue
            t0 = time.perf_counter()
            sheet = _extract_sheet(
                wb_f[part.name], wb_v[part.name], part, index, info.epoch, inventory
            )
            log.info(
                "extracted sheet %r: %d cells, %d formulas in %.1fs",
                sheet.name,
                sheet.cell_count,
                sheet.formula_count,
                time.perf_counter() - t0,
            )
            if on_sheet is not None:
                on_sheet(sheet)
            if not keep_cells:
                sheet = sheet.model_copy(update={"cells": CellColumns()})
            sheets.append(sheet)
    finally:
        wb_f.close()
        wb_v.close()

    log.info("parsed %s in %.1fs", path.name, time.perf_counter() - started)
    return RawWorkbook(
        filename=filename or path.name,
        epoch=1904 if info.epoch == 1904 else 1900,
        has_macros=info.has_macros or suffix == ".xlsm",
        warnings=warnings,
        defined_names=info.defined_names,
        tables=[t for part in info.sheets for t in part.tables],
        functions=inventory.as_dict(),
        sheets=sheets,
        pivots=pivots,
    )
