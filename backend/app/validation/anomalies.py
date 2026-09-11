"""Structural anomalies: warnings a researcher can act on, computed from the logic model and raw cells.

- fragmentation: a short block interrupting a uniform run of one template in a column
- own_row: a short block whose relative row offsets differ from the surrounding pattern
  (the cut-and-paste signature: the row reads another row's data)
- duplicate_keys: repeated keys in exact-match lookup ranges (first match wins silently)
- stale: formula cells saved without a cached value (no full recalculation before saving)
"""

from __future__ import annotations

from collections import defaultdict
from collections.abc import Callable

from openpyxl.utils import get_column_letter

from app.model.formula.ast import Call, CellRef, RangeRef, iter_nodes, iter_refs, ref_rect
from app.model.formula.refs import Rect, a1_cell
from app.model.schema import FormulaBlock, Template, WorkbookLogicModel
from app.parser.models import RawSheet
from app.validation.schema import Anomaly

SHORT_ROWS = 2
LARGE_ROWS = 20
MAX_DUP_KEYS = 50
LOOKUPS = {"VLOOKUP": 0, "HLOOKUP": 1}

RawLoader = Callable[[str], RawSheet]


def _row_offsets(tpl: Template) -> set[int]:
    if tpl.ast is None:
        return set()
    offsets: set[int] = set()
    for ref in iter_refs(tpl.ast):
        if isinstance(ref, CellRef) and not ref.row.abs:
            offsets.add(ref.row.v)
        elif isinstance(ref, RangeRef):
            for axis in (ref.r1, ref.r2):
                if not axis.abs:
                    offsets.add(axis.v)
    return offsets


def _columns_blocks(model: WorkbookLogicModel) -> dict[tuple[str, int], list[FormulaBlock]]:
    """(sheet, column) -> formula blocks covering that column, sorted by row."""
    out: dict[tuple[str, int], list[FormulaBlock]] = defaultdict(list)
    for b in model.formula_blocks:
        r = b.rect.rect()
        for c in range(r.c1, r.c2 + 1):
            out[(b.sheet, c)].append(b)
    for lst in out.values():
        lst.sort(key=lambda b: b.rect.r1)
    return out


def fragmentation_and_own_row(
    model: WorkbookLogicModel, next_id: Callable[[], int]
) -> list[Anomaly]:
    columns = _columns_blocks(model)
    frag_rows: dict[tuple[str, int, int, str], set[int]] = defaultdict(
        set
    )  # (sheet, r1, r2, kind) -> columns
    details: dict[tuple[str, int, int, str], dict] = {}
    for (sheet, col), blocks in columns.items():
        for i, b in enumerate(blocks):
            rect = b.rect.rect()
            if rect.rows > SHORT_ROWS:
                continue
            above = blocks[i - 1] if i > 0 and blocks[i - 1].rect.r2 == rect.r1 - 1 else None
            below = (
                blocks[i + 1]
                if i + 1 < len(blocks) and blocks[i + 1].rect.r1 == rect.r2 + 1
                else None
            )
            big_above = above is not None and above.rect.rect().rows >= LARGE_ROWS
            big_below = below is not None and below.rect.rect().rows >= LARGE_ROWS
            tpl = model.template(b.template_id)
            neighbour = above if big_above else below if big_below else None
            if neighbour is None or neighbour.template_id == b.template_id:
                continue
            ntpl = model.template(neighbour.template_id)
            # Fragmentation: the fragment interrupts one template running above and below it.
            if big_above and big_below and above.template_id == below.template_id:
                key = (sheet, rect.r1, rect.r2, "fragmentation")
                frag_rows[key].add(col)
                details.setdefault(
                    key,
                    {
                        "fragment_formula": tpl.example_formula,
                        "neighbour_formula": ntpl.example_formula,
                        "neighbour_rows": f"{above.rect.r1}–{below.rect.r2}",
                    },
                )
            # Own-row violation: relative row offsets the neighbouring pattern does not use.
            own, theirs = _row_offsets(tpl), _row_offsets(ntpl)
            foreign = {o for o in own - theirs if o != 0}
            if foreign and all(o == 0 for o in theirs):
                key = (sheet, rect.r1, rect.r2, "own_row")
                frag_rows[key].add(col)
                target_rows = sorted({rect.r1 + o for o in foreign})
                details.setdefault(
                    key,
                    {
                        "fragment_formula": tpl.example_formula,
                        "neighbour_formula": ntpl.example_formula,
                        "reads_rows": target_rows,
                        "offsets": sorted(foreign),
                    },
                )
    anomalies: list[Anomaly] = []
    for (sheet, r1, r2, kind), cols in sorted(frag_rows.items()):
        col_list = sorted(cols)
        letters = [get_column_letter(c) for c in col_list]
        location = ", ".join(
            f"{letter}{r1}" if r1 == r2 else f"{letter}{r1}:{letter}{r2}" for letter in letters
        )
        rows = f"row {r1}" if r1 == r2 else f"rows {r1}–{r2}"
        d = details[(sheet, r1, r2, kind)]
        if kind == "fragmentation":
            title = f"{sheet} {rows}: {len(col_list)} column(s) break the surrounding pattern"
            explanation = (
                f"Columns {', '.join(letters)} on {rows} use a different formula from the rows above and below "
                f"(neighbours: {d['neighbour_formula'][:80]} · here: {d['fragment_formula'][:80]}). "
                "A row that was pasted in, or edited by hand, looks exactly like this. Check it against its neighbours."
            )
        else:
            reads = ", ".join(str(r) for r in d["reads_rows"][:4])
            title = f"{sheet} {rows}: formulas read row {reads} instead of their own row"
            explanation = (
                f"Columns {', '.join(letters)} on {rows} reference row {reads} while the surrounding rows reference "
                f"themselves ({d['neighbour_formula'][:80]}). This is the signature of a cut-and-pasted row: the "
                "label on this row does not describe the data its formulas use. Re-enter the formulas so they point at this row."
            )
        anomalies.append(
            Anomaly(
                id=next_id(),
                kind=kind,
                sheet=sheet,
                location=location,
                title=title,
                explanation=explanation,
                detail={**d, "columns": letters},
            )  # type: ignore[arg-type]
        )
    return anomalies


def duplicate_keys(
    model: WorkbookLogicModel, load_raw: RawLoader, next_id: Callable[[], int]
) -> list[Anomaly]:
    """Repeated keys in the first column/row of exact-match VLOOKUP/HLOOKUP ranges."""
    wanted: dict[tuple[str, str, int], set[int]] = defaultdict(
        set
    )  # (sheet, range a1, axis) -> template ids
    extents = {s.name: Rect.from_a1(s.used_range) for s in model.sheets if s.used_range}
    for tpl in model.templates:
        if tpl.ast is None:
            continue
        r, c = _example_rc(tpl)
        for node in iter_nodes(tpl.ast):
            if not (isinstance(node, Call) and node.name in LOOKUPS and len(node.args) >= 4):
                continue
            flag = node.args[3]
            if not (getattr(flag, "kind", "") == "bool" and flag.value is False) and not (
                getattr(flag, "kind", "") == "num" and flag.value == 0
            ):
                continue
            rng = node.args[1]
            if not isinstance(rng, RangeRef | CellRef) and type(rng).__name__ not in (
                "ColumnRef",
                "RowRef",
            ):
                continue
            sheet = getattr(rng, "sheet", None) or tpl.sheet
            rect = ref_rect(rng, r, c, extents.get(sheet))  # type: ignore[arg-type]
            if rect is None:
                continue
            wanted[(sheet, rect.to_a1(), LOOKUPS[node.name])].add(tpl.id)
    anomalies: list[Anomaly] = []
    cache: dict[str, RawSheet] = {}
    for (sheet, a1, axis), tpl_ids in sorted(wanted.items()):
        raw = cache.get(sheet)
        if raw is None:
            try:
                raw = load_raw(sheet)
            except Exception:  # noqa: BLE001 - a missing sheet is not an anomaly
                continue
            cache[sheet] = raw
        rect = Rect.from_a1(a1)
        keys: dict[str, list[str]] = defaultdict(list)
        cells = rect.rows if axis == 0 else rect.cols
        for k in range(cells):
            addr = a1_cell(rect.r1 + k, rect.c1) if axis == 0 else a1_cell(rect.r1, rect.c1 + k)
            cell = raw.cell(addr)
            if cell is None or cell.value is None or cell.value_type in ("empty", "error"):
                continue
            key = cell.value.casefold() if isinstance(cell.value, str) else repr(cell.value)
            keys[key].append(addr)
        # Header labels repeated across header rows are not lookup keys. For column lookups
        # (VLOOKUP) keep a duplicate only when some occurrence sits inside the sheet's data body
        # (a block spanning at least LARGE_ROWS rows); HLOOKUP keys live in header rows by nature.
        data_rows = _data_rows(model, sheet) if axis == 0 else None
        dups = {
            k: v
            for k, v in keys.items()
            if len(v) > 1 and (data_rows is None or any(_row_of(addr) in data_rows for addr in v))
        }
        if not dups:
            continue
        sample = list(dups.items())[:MAX_DUP_KEYS]
        readers = ", ".join(
            f"{model.template(t).sheet}!{model.template(t).example_cell}"
            for t in sorted(tpl_ids)[:3]
        )
        anomalies.append(
            Anomaly(
                id=next_id(),
                kind="duplicate_keys",
                sheet=sheet,
                location=a1,
                title=f"{sheet}!{a1}: {len(dups)} duplicate lookup key(s)",
                explanation=(
                    f"Exact-match lookups into {sheet}!{a1} (from {readers}) return the first row that matches. "
                    f"{len(dups)} key(s) appear more than once, so later duplicates are silently ignored; "
                    "if they carry different data, the workbook picks one without telling you. First duplicates: "
                    + "; ".join(
                        f"{raw.cell(v[0]).value!r} at {', '.join(v[:4])}" for _, v in sample[:3]
                    )
                ),
                detail={
                    "duplicates": [
                        {"key": raw.cell(v[0]).value, "cells": v[:10], "count": len(v)}
                        for _, v in sample
                    ],
                    "readers": sorted(tpl_ids),
                },
            )
        )
    return anomalies


def stale_cells(
    model: WorkbookLogicModel, load_raw: RawLoader, next_id: Callable[[], int]
) -> list[Anomaly]:
    anomalies: list[Anomaly] = []
    for s in model.sheets:
        if not s.in_scope or s.formula_cells == 0:
            continue
        raw = load_raw(s.name)
        stale = [
            a
            for a, f, t in zip(
                raw.cells.address, raw.cells.formula, raw.cells.value_type, strict=True
            )
            if f and t == "empty"
        ]
        if not stale:
            continue
        anomalies.append(
            Anomaly(
                id=next_id(),
                kind="stale",
                sheet=s.name,
                location=", ".join(stale[:5]) + (", …" if len(stale) > 5 else ""),
                title=f"{s.name}: {len(stale)} formula cell(s) have no cached value",
                explanation=(
                    "Excel saved these formulas without a computed result, which happens when a workbook is saved "
                    "with calculation set to manual or before a full recalculation. They cannot be reconciled; open the "
                    "master, press F9 (or set calculation to automatic), save, and upload again."
                ),
                detail={"count": len(stale), "cells": stale[:50]},
            )
        )
    return anomalies


def find_anomalies(model: WorkbookLogicModel, load_raw: RawLoader) -> list[Anomaly]:
    counter = iter(range(1, 1_000_000))

    def next_id() -> int:
        return next(counter)

    out = fragmentation_and_own_row(model, next_id)
    out += duplicate_keys(model, load_raw, next_id)
    out += stale_cells(model, load_raw, next_id)
    return out


def _example_rc(tpl: Template) -> tuple[int, int]:
    from app.model.formula.refs import parse_a1_cell

    return parse_a1_cell(tpl.example_cell)


def _data_rows(model: WorkbookLogicModel, sheet: str) -> set[int]:
    """Rows of ``sheet`` covered by any block (input or formula) spanning at least LARGE_ROWS rows."""
    rows: set[int] = set()
    for block in [*model.input_blocks, *model.formula_blocks]:
        if block.sheet != sheet:
            continue
        r = block.rect.rect()
        if r.rows >= LARGE_ROWS:
            rows.update(range(r.r1, r.r2 + 1))
    return rows


def _row_of(addr: str) -> int:
    from app.model.formula.refs import parse_a1_cell

    return parse_a1_cell(addr)[0]
