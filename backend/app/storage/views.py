"""Read models for the dashboard views: a windowed sheet grid, the inputs catalogue and the
outputs summary.

Everything here is derived from the logic model, the raw sheet and a run's values. Labels come
from the model's label detection or ``dashboard.config.json`` ``labelOverrides``; nothing in
this module knows any particular workbook.
"""

from __future__ import annotations

import math
from collections import Counter
from typing import Any

from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import from_excel
from pydantic import BaseModel, Field
from sqlalchemy.orm import Session

from app.dashboard_config import DashboardConfig
from app.model.formula.refs import Rect, a1_cell, parse_a1_cell
from app.model.schema import BusinessRule, FormulaBlock, InputBlock, WorkbookLogicModel
from app.parser.models import RawSheet
from app.storage import runs, workbooks
from app.storage.models import Run

MAX_ROWS = 200
MAX_COLS = 120
DEFAULT_ROWS = 100
PARAMETER_CELLS = 64  # an input block this small ...
PARAMETER_ROWS = 8  # ... and this short is a set of parameters, not a table
METRIC_CELLS = 12  # an output block this small is a set of headline numbers
BODY_MIN_ROWS = 20  # a block this tall marks where a sheet's body starts
SERIES_MIN_ROWS = 50
SERIES_MAX_POINTS = 400
SERIES_MAX_COLUMNS = 4
LABEL_ROWS = 5
LABEL_COLS = 3
RULES_PER_SHEET = 50


# ---- raw sheet access ------------------------------------------------------------------------


class SheetIndex:
    """A raw sheet with (row, col) addressing."""

    def __init__(self, raw: RawSheet) -> None:
        self.raw = raw
        c = raw.cells
        self.at: dict[tuple[int, int], int] = {
            (r, col): i for i, (r, col) in enumerate(zip(c.row, c.col, strict=True))
        }
        self.used = Rect.from_a1(raw.used_range) if raw.used_range else None

    def value(self, row: int, col: int) -> Any:
        i = self.at.get((row, col))
        return None if i is None else self.raw.cells.value[i]

    def value_type(self, row: int, col: int) -> str:
        i = self.at.get((row, col))
        return "empty" if i is None else self.raw.cells.value_type[i]

    def formula(self, row: int, col: int) -> str | None:
        i = self.at.get((row, col))
        return None if i is None else self.raw.cells.formula[i]

    def number_format(self, row: int, col: int) -> str | None:
        i = self.at.get((row, col))
        return None if i is None else self.raw.cells.number_format[i]

    def text(self, row: int, col: int) -> str | None:
        v = self.value(row, col)
        if isinstance(v, str) and v.strip() and self.formula(row, col) is None:
            return v.strip()
        return None

    def any_text(self, row: int, col: int) -> str | None:
        v = self.value(row, col)
        return v.strip() if isinstance(v, str) and v.strip() else None


class _SheetCache:
    def __init__(self, capacity: int = 20) -> None:
        self.capacity = capacity
        self._items: dict[tuple[str, str], SheetIndex] = {}

    def get(self, session: Session, version_id: str, sheet: str) -> SheetIndex:
        key = (version_id, sheet)
        hit = self._items.get(key)
        if hit is None:
            raw = workbooks.load_raw(session, version_id, sheet=sheet).sheets[0]
            hit = SheetIndex(raw)
            self._items[key] = hit
            while len(self._items) > self.capacity:
                self._items.pop(next(iter(self._items)))
        return hit

    def clear(self) -> None:
        self._items.clear()


sheet_cache = _SheetCache()


# ---- shared helpers --------------------------------------------------------------------------


def _root_run(session: Session, run: Run) -> Run:
    current = run
    while current.parent_run_id:
        parent = session.get(Run, current.parent_run_id)
        if parent is None:
            break
        current = parent
    return current


def _overrides_for(session: Session, run: Run | None, sheet: str) -> dict[str, Any]:
    """Override values on one sheet, address -> value, composed over the run's parent."""
    if run is None:
        return {}
    merged: dict[str, Any] = {}
    chain: list[Run] = []
    current: Run | None = run
    while current is not None:
        chain.append(current)
        current = session.get(Run, current.parent_run_id) if current.parent_run_id else None
    for r in reversed(chain):
        merged.update(runs.overrides_of(r))
    out: dict[str, Any] = {}
    for key, value in merged.items():
        if "!" in key and key.rsplit("!", 1)[0].strip("'") == sheet:
            out[key.rsplit("!", 1)[1]] = value
    return out


def _python_type(v: Any) -> str:
    if v is None:
        return "empty"
    if isinstance(v, bool):
        return "bool"
    if isinstance(v, int | float):
        return "number"
    return "text"


def _blocks_on(
    model: WorkbookLogicModel, sheet: str
) -> tuple[list[FormulaBlock], list[InputBlock]]:
    return (
        [b for b in model.formula_blocks if b.sheet == sheet],
        [b for b in model.input_blocks if b.sheet == sheet],
    )


def _body_start(idx: SheetIndex, fblocks: list[FormulaBlock], iblocks: list[InputBlock]) -> int:
    tall = [b.rect.r1 for b in [*fblocks, *iblocks] if b.rect.rect().rows >= BODY_MIN_ROWS]
    if tall:
        return min(tall)
    return idx.used.r1 if idx.used else 1


def _column_label(
    col: int, fblocks: list[FormulaBlock], iblocks: list[InputBlock], idx: SheetIndex, body: int
) -> str | None:
    for block in [*fblocks, *iblocks]:
        r = block.rect
        if r.c1 <= col <= r.c2 and r.r2 >= body:
            label = (
                block.column_labels[col - r.c1] if col - r.c1 < len(block.column_labels) else None
            )
            if label:
                return label
    for row in range(body - 1, max(0, body - 1 - LABEL_ROWS), -1):
        text = idx.any_text(row, col)
        if text:
            return text
    return None


def _nearest_label(idx: SheetIndex, row: int, col: int) -> str | None:
    for c in range(col - 1, max(0, col - 1 - LABEL_COLS), -1):
        text = idx.text(row, c)
        if text:
            return text
    for r in range(row - 1, max(0, row - 1 - LABEL_ROWS), -1):
        text = idx.any_text(r, col)
        if text:
            return text
    return None


def _cell_label(
    cfg: DashboardConfig,
    sheet: str,
    row: int,
    col: int,
    block: FormulaBlock | InputBlock,
    idx: SheetIndex,
) -> str | None:
    override = cfg.label_overrides.get(f"{sheet}!{a1_cell(row, col)}")
    if override:
        return override
    if block.row_label_col is not None:
        text = idx.any_text(row, block.row_label_col)
        if text:
            return text
    i = col - block.rect.c1
    if 0 <= i < len(block.column_labels) and block.column_labels[i]:
        return block.column_labels[i]
    return _nearest_label(idx, row, col)


def _format_kind(fmt: str | None, vtype: str) -> str:
    if vtype == "date":
        return "date"
    if not fmt or fmt == "General":
        return "general"
    f = fmt.lower()
    if "%" in f:
        return "percent"
    if any(token in f for token in ("yy", "mmm", "dd", "d-m", "m/d", "d/m")):
        return "date"
    if "." in f:
        return "number"
    if "0" in f:
        return "integer"
    return "general"


def _section_label(idx: SheetIndex, rect: Rect) -> str | None:
    for row in range(rect.r1 - 1, max(0, rect.r1 - 1 - LABEL_ROWS), -1):
        for col in range(rect.c1, min(rect.c2, rect.c1 + LABEL_COLS) + 1):
            text = idx.any_text(row, col)
            if text:
                return text
    for col in range(rect.c1 - 1, max(0, rect.c1 - 1 - LABEL_COLS), -1):
        text = idx.text(rect.r1, col)
        if text:
            return text
    return None


# ---- grid window -----------------------------------------------------------------------------


class GridColumn(BaseModel):
    col: int
    letter: str
    label: str | None = None
    kind: str = "static"  # output | formula | input | static
    format: str = "general"  # general | number | integer | percent | date | text


class GridRow(BaseModel):
    row: int
    label: str | None = None
    cells: list[dict[str, Any] | None]


class GridOut(BaseModel):
    run_id: str
    sheet: str
    r1: int
    r2: int
    c1: int
    c2: int
    used_range: str | None
    body_start: int
    total_rows: int
    total_cols: int
    row_label_col: int | None = None
    baseline_run_id: str | None = None
    columns: list[GridColumn]
    rows: list[GridRow]


def grid_window(
    session: Session,
    run: Run,
    model: WorkbookLogicModel,
    sheet: str,
    *,
    r1: int | None = None,
    r2: int | None = None,
    c1: int | None = None,
    c2: int | None = None,
) -> GridOut:
    idx = sheet_cache.get(session, run.version_id, sheet)
    fblocks, iblocks = _blocks_on(model, sheet)
    used = idx.used
    if used is None:
        return GridOut(
            run_id=run.id, sheet=sheet, r1=1, r2=0, c1=1, c2=0, used_range=None, body_start=1,
            total_rows=0, total_cols=0, columns=[], rows=[],
        )  # fmt: skip
    body = _body_start(idx, fblocks, iblocks)
    r1 = max(used.r1, r1 if r1 is not None else body)
    r2 = min(used.r2, r2 if r2 is not None else r1 + DEFAULT_ROWS - 1, r1 + MAX_ROWS - 1)
    c1 = max(used.c1, c1 if c1 is not None else used.c1)
    c2 = min(used.c2, c2 if c2 is not None else used.c2, c1 + MAX_COLS - 1)

    computed = runs.formula_values(session, run, sheet)
    root = _root_run(session, run)
    baseline = runs.formula_values(session, root, sheet) if root.id != run.id else {}
    overrides = _overrides_for(session, run, sheet)
    output_cols = {
        c for b in fblocks if b.classification == "output" for c in range(b.rect.c1, b.rect.c2 + 1)
    }
    formula_cols = {c for b in fblocks for c in range(b.rect.c1, b.rect.c2 + 1)}
    input_cols = {c for b in iblocks for c in range(b.rect.c1, b.rect.c2 + 1)}

    columns: list[GridColumn] = []
    for col in range(c1, c2 + 1):
        kind = (
            "output"
            if col in output_cols
            else "formula"
            if col in formula_cols
            else "input"
            if col in input_cols
            else "static"
        )
        fmt = "general"
        for row in range(max(body, r1), min(r2, max(body, r1) + 5) + 1):
            if (row, col) in idx.at:
                fmt = _format_kind(idx.number_format(row, col), idx.value_type(row, col))
                break
        columns.append(
            GridColumn(
                col=col,
                letter=get_column_letter(col),
                label=_column_label(col, fblocks, iblocks, idx, body),
                kind=kind,
                format=fmt,
            )
        )

    label_cols = Counter(
        b.row_label_col
        for b in fblocks
        if b.row_label_col is not None and b.rect.rect().intersects(Rect(r1, c1, r2, c2))
    )
    row_label_col = label_cols.most_common(1)[0][0] if label_cols else None
    if row_label_col is None:
        for col in range(used.c1, min(used.c2, used.c1 + LABEL_COLS) + 1):
            sample = [idx.text(row, col) for row in range(body, min(used.r2, body + 40) + 1)]
            if sample and sum(1 for s in sample if s) >= 0.6 * len(sample):
                row_label_col = col
                break

    rows: list[GridRow] = []
    for row in range(r1, r2 + 1):
        cells: list[dict[str, Any] | None] = []
        for col in range(c1, c2 + 1):
            i = idx.at.get((row, col))
            if i is None:
                cells.append(None)
                continue
            addr = idx.raw.cells.address[i]
            if idx.raw.cells.formula[i] is not None:
                if addr in computed:
                    v, t = computed[addr]
                    cell: dict[str, Any] = {"v": v, "t": t, "f": True}
                else:
                    cell = {
                        "v": idx.raw.cells.value[i],
                        "t": idx.raw.cells.value_type[i],
                        "f": True,
                        "stale": True,
                    }
                if addr in baseline and baseline[addr][0] != cell["v"]:
                    cell["b"] = baseline[addr][0]
            elif addr in overrides:
                cell = {"v": overrides[addr], "t": _python_type(overrides[addr]), "o": True}
            else:
                cell = {"v": idx.raw.cells.value[i], "t": idx.raw.cells.value_type[i]}
            cells.append(cell)
        label = idx.any_text(row, row_label_col) if row_label_col is not None else None
        rows.append(GridRow(row=row, label=label, cells=cells))

    return GridOut(
        run_id=run.id,
        sheet=sheet,
        r1=r1,
        r2=r2,
        c1=c1,
        c2=c2,
        used_range=used.to_a1(),
        body_start=body,
        total_rows=used.r2,
        total_cols=used.c2,
        row_label_col=row_label_col,
        baseline_run_id=root.id if root.id != run.id else None,
        columns=columns,
        rows=rows,
    )


# ---- inputs catalogue ------------------------------------------------------------------------


def _input_shape(block: InputBlock, rect: Rect) -> str:
    """parameters: a few labelled cells (the what-if levers); labels: a header row or column of
    text that lookups scan (read, rarely edited); table: everything larger."""
    if block.cell_count > PARAMETER_CELLS or rect.rows > PARAMETER_ROWS:
        return "table"
    text = block.value_types.get("text", 0)
    if text >= 0.8 * max(block.cell_count, 1) and (rect.rows == 1 or rect.cols == 1):
        return "labels"
    return "parameters"


class InputCellOut(BaseModel):
    address: str
    label: str | None = None
    value: Any = None
    type: str = "empty"
    override: bool = False


class InputBlockOut(BaseModel):
    id: int
    sheet: str
    rect: str
    r1: int
    c1: int
    r2: int
    c2: int
    cell_count: int
    kind: str  # input | external
    editable: bool
    shape: str  # parameters | labels | table
    section: str | None = None
    value_types: dict[str, int] = Field(default_factory=dict)
    column_labels: list[str | None] = Field(default_factory=list)
    row_label_col: int | None = None
    cells: list[InputCellOut] = Field(default_factory=list)


class InputSheetOut(BaseModel):
    sheet: str
    role: str
    role_source: str
    blocks: list[InputBlockOut]


class InputsOut(BaseModel):
    run_id: str | None
    sheets: list[InputSheetOut]


def inputs_catalogue(
    session: Session,
    model: WorkbookLogicModel,
    version_id: str,
    run: Run | None,
    cfg: DashboardConfig,
) -> InputsOut:
    sheets: list[InputSheetOut] = []
    for sm in model.sheets:
        if not sm.in_scope:
            continue
        iblocks = sorted(
            (b for b in model.input_blocks if b.sheet == sm.name),
            key=lambda b: (b.rect.r1, b.rect.c1),
        )
        if not iblocks:
            continue
        idx = sheet_cache.get(session, version_id, sm.name)
        overrides = _overrides_for(session, run, sm.name)
        out_blocks: list[InputBlockOut] = []
        for b in iblocks:
            rect = b.rect.rect()
            shape = _input_shape(b, rect)
            cells: list[InputCellOut] = []
            if shape in ("parameters", "labels"):
                for row in range(rect.r1, rect.r2 + 1):
                    for col in range(rect.c1, rect.c2 + 1):
                        i = idx.at.get((row, col))
                        if i is None or idx.raw.cells.formula[i] is not None:
                            continue
                        addr = idx.raw.cells.address[i]
                        if addr in overrides:
                            value, vtype, ov = overrides[addr], _python_type(overrides[addr]), True
                        else:
                            value, vtype, ov = (
                                idx.raw.cells.value[i],
                                idx.raw.cells.value_type[i],
                                False,
                            )
                        cells.append(
                            InputCellOut(
                                address=addr,
                                label=_cell_label(cfg, sm.name, row, col, b, idx),
                                value=value,
                                type=vtype,
                                override=ov,
                            )
                        )
            out_blocks.append(
                InputBlockOut(
                    id=b.id,
                    sheet=b.sheet,
                    rect=rect.to_a1(),
                    r1=rect.r1,
                    c1=rect.c1,
                    r2=rect.r2,
                    c2=rect.c2,
                    cell_count=b.cell_count,
                    kind=b.kind,
                    editable=b.kind == "input",
                    shape=shape,
                    section=cfg.label_overrides.get(f"{sm.name}!{rect.to_a1()}")
                    or _section_label(idx, rect),
                    value_types=b.value_types,
                    column_labels=b.column_labels,
                    row_label_col=b.row_label_col,
                    cells=cells,
                )
            )
        sheets.append(
            InputSheetOut(
                sheet=sm.name, role=sm.role, role_source=sm.role_source, blocks=out_blocks
            )
        )
    return InputsOut(run_id=run.id if run else None, sheets=sheets)


# ---- outputs summary -------------------------------------------------------------------------


class MetricOut(BaseModel):
    address: str
    label: str | None = None
    value: Any = None
    type: str = "empty"
    format: str = "general"
    baseline: Any = None
    delta: float | None = None


class SeriesColumn(BaseModel):
    label: str
    points: list[list[Any]]  # [x, y]


class SeriesOut(BaseModel):
    x_label: str | None
    x_type: str  # date | number
    rows: int
    columns: list[SeriesColumn]


class OutputBlockOut(BaseModel):
    id: int
    rect: str
    cell_count: int
    kind: str  # metric | table
    labels: list[str | None] = Field(default_factory=list)


class OutputSheetOut(BaseModel):
    sheet: str
    role: str
    role_source: str
    body_start: int
    table_rect: str | None = None
    blocks: list[OutputBlockOut]
    metrics: list[MetricOut]
    rules: list[BusinessRule]
    series: SeriesOut | None = None
    changed_cells: int = 0


class OutputsOut(BaseModel):
    run_id: str
    baseline_run_id: str | None
    sheets: list[OutputSheetOut]


def output_sheet_order(model: WorkbookLogicModel, cfg: DashboardConfig) -> list[str]:
    in_scope = {s.name: s for s in model.sheets if s.in_scope}
    ordered = [s for s in cfg.output_sheets if s in in_scope]
    ordered += [
        s.name for s in model.sheets if s.in_scope and s.role == "output" and s.name not in ordered
    ]
    return ordered


def _numeric(v: Any) -> bool:
    return isinstance(v, int | float) and not isinstance(v, bool) and math.isfinite(float(v))


def _series(
    idx: SheetIndex, computed: dict[str, tuple[Any, str]], fblocks: list[FormulaBlock], body: int
) -> SeriesOut | None:
    if idx.used is None:
        return None
    c = idx.raw.cells
    date_cols: Counter[int] = Counter()
    for i in range(len(c)):
        if c.row[i] >= body and c.value_type[i] == "date" and c.formula[i] is None:
            date_cols[c.col[i]] += 1
    if not date_cols:
        return None
    x_col, n_dates = date_cols.most_common(1)[0]
    if n_dates < SERIES_MIN_ROWS:
        return None
    numeric_cols: Counter[int] = Counter()
    by_addr: dict[tuple[int, int], Any] = {}
    for addr, (v, _t) in computed.items():
        if _numeric(v):
            row, col = parse_a1_cell(addr)
            if row >= body:
                numeric_cols[col] += 1
                by_addr[(row, col)] = v
    cols = [col for col, n in sorted(numeric_cols.items()) if n >= SERIES_MIN_ROWS][
        :SERIES_MAX_COLUMNS
    ]
    if not cols:
        return None
    rows = sorted(r for (r, col), _ in idx.at.items() if col == x_col and r >= body)
    stride = max(1, math.ceil(len(rows) / SERIES_MAX_POINTS))
    sampled = rows[::stride]

    def x_of(row: int) -> Any:
        v = idx.value(row, x_col)
        try:
            return from_excel(float(v)).date().isoformat()
        except (TypeError, ValueError, OverflowError):
            return v

    series_cols: list[SeriesColumn] = []
    for col in cols:
        label = _column_label(col, fblocks, [], idx, body) or get_column_letter(col)
        points = [[x_of(r), by_addr.get((r, col))] for r in sampled]
        series_cols.append(SeriesColumn(label=label, points=points))
    return SeriesOut(
        x_label=_column_label(x_col, fblocks, [], idx, body),
        x_type="date",
        rows=len(rows),
        columns=series_cols,
    )


def outputs_summary(
    session: Session, model: WorkbookLogicModel, run: Run, cfg: DashboardConfig
) -> OutputsOut:
    root = _root_run(session, run)
    has_baseline = root.id != run.id
    sheets: list[OutputSheetOut] = []
    for name in output_sheet_order(model, cfg):
        sm = model.sheet(name)
        assert sm is not None
        idx = sheet_cache.get(session, run.version_id, name)
        fblocks, iblocks = _blocks_on(model, name)
        body = _body_start(idx, fblocks, iblocks)
        computed = runs.formula_values(session, run, name)
        baseline = runs.formula_values(session, root, name) if has_baseline else {}
        out_blocks = [b for b in fblocks if b.classification == "output"]
        blocks: list[OutputBlockOut] = []
        metrics: list[MetricOut] = []
        table_rect: Rect | None = None
        for b in sorted(out_blocks, key=lambda b: (b.rect.r1, b.rect.c1)):
            rect = b.rect.rect()
            kind = "metric" if b.cell_count <= METRIC_CELLS else "table"
            blocks.append(
                OutputBlockOut(
                    id=b.id,
                    rect=rect.to_a1(),
                    cell_count=b.cell_count,
                    kind=kind,
                    labels=b.column_labels,
                )
            )
            if kind == "table":
                table_rect = rect if table_rect is None else table_rect.bbox(rect)
                continue
            for row in range(rect.r1, rect.r2 + 1):
                for col in range(rect.c1, rect.c2 + 1):
                    addr = a1_cell(row, col)
                    v, t = computed.get(addr, (idx.value(row, col), idx.value_type(row, col)))
                    m = MetricOut(
                        address=addr,
                        label=_cell_label(cfg, name, row, col, b, idx),
                        value=v,
                        type=t,
                    )
                    if addr in baseline:
                        bv = baseline[addr][0]
                        m.baseline = bv
                        if _numeric(v) and _numeric(bv) and v != bv:
                            m.delta = float(v) - float(bv)
                    metrics.append(m)
        changed = 0
        if has_baseline:
            changed = sum(
                1
                for addr, (v, _t) in computed.items()
                if addr in baseline and baseline[addr][0] != v
            )
        sheets.append(
            OutputSheetOut(
                sheet=name,
                role=sm.role,
                role_source=sm.role_source,
                body_start=body,
                table_rect=table_rect.to_a1() if table_rect else None,
                blocks=blocks,
                metrics=metrics,
                rules=[r for r in model.rules if r.sheet == name][:RULES_PER_SHEET],
                series=_series(idx, computed, fblocks, body),
                changed_cells=changed,
            )
        )
    return OutputsOut(
        run_id=run.id, baseline_run_id=root.id if has_baseline else None, sheets=sheets
    )
