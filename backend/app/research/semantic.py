# ruff: noqa: N815  -- field names mirror the JSON keys researchers edit
"""The semantic map: what a row means, declared in ``dashboard.config.json`` under ``research``.

Nothing here knows a sheet name; the map is data. ``resolve_map`` checks every reference
against a version's logic model and raw sheets and returns readable problems instead of
raising, so an incomplete map degrades to a "not configured" state rather than a 500.
"""

from __future__ import annotations

import re
from collections.abc import Callable
from difflib import get_close_matches
from typing import Any, Literal

from openpyxl.utils import column_index_from_string, get_column_letter
from pydantic import BaseModel, Field, ValidationError

from app.model.schema import WorkbookLogicModel

Role = Literal["score", "rank", "quartile", "return", "factor"]
Op = Literal["eq", "ne", "lt", "lte", "gt", "gte", "in", "top", "bottom", "notnull", "isnull"]
# {name} · {name|one|many} renders "N one/many" · {name?one|many} renders just the word.
SENTENCE_PLACEHOLDER = re.compile(
    r"\{([a-z_][a-z0-9_]*(?:\.[a-z_][a-z0-9_]*)*)(?:([|?])([^{}|]*)\|([^{}|]*))?\}"
)


def placeholder_names(template: str) -> list[str]:
    return [m.group(1) for m in SENTENCE_PLACEHOLDER.finditer(template)]


INSIGHT_PLACEHOLDERS = {
    "count", "total", "pct", "sum", "versions", "groups", "min_group",
    "top.label", "top.value", "top.sub", "top.group",
    "group.label", "group.value", "group.count", "group.total",
    "hit.label", "hit.value", "hit.count", "hit.total", "measure.label",
}  # fmt: skip
NARRATIVE_PLACEHOLDERS: dict[str, set[str]] = {
    "universe": {"total", "rated", "unrated", "categories", "unranked", "unranked_below", "unrated_small", "unrated_other", "unrated_young", "unrated_gap", "unrated_unknown", "coverage_since", "outside", "as_of"},
    "dashboard": {"moved", "up", "down", "repairs", "held_q1", "versions", "previous", "current", "rated", "categories"},
    "movement": {"rated", "moved", "up", "down", "repairs", "into_q1", "out_of_q1", "entries", "exits", "previous", "current"},
    "categories": {"categories", "unranked", "unranked_below", "eligible", "min_group", "widest.label", "widest.value", "best.label", "best.value", "measure.label"},
    "fund": {"label", "rank", "category_count", "category", "quartile", "delta", "delta_text", "previous", "score"},
    "executive": {"rated", "total", "q1", "q1_pct", "ranked_categories", "categories", "unranked", "unranked_below", "unrated", "unrated_young", "unrated_small", "unrated_gap", "unrated_unknown", "outside", "moved", "up", "down", "repairs", "previous", "current", "checked", "matched", "mismatched", "agreement", "scope", "as_of"},
    "distribution": {"rated", "ranked_categories", "categories", "median_category", "largest_category", "largest_category_rated"},
    "coverage": {"total", "rated", "unrated", "categories", "unranked", "unranked_below", "unrated_small", "unrated_other", "unrated_young", "unrated_gap", "unrated_unknown", "coverage_since", "outside", "as_of"},
}  # fmt: skip
# Executive sentences and KPI tiles may also read a computed insight: {ins.<key>.count},
# {ins.<key>.top.label}, {ins.<key>.group.count} ... (any placeholder an insight sentence has).
INSIGHT_REF = re.compile(r"^ins\.([a-z_][a-z0-9_]*)\.(.+)$")


def narrative_placeholder_ok(name: str, ph: str, rmap: ResearchMap) -> bool:
    allowed = NARRATIVE_PLACEHOLDERS.get(name)
    if allowed is None:
        return False
    if ph in allowed:
        return True
    if name in ("executive", "kpi"):
        m = INSIGHT_REF.match(ph)
        return (
            bool(m)
            and any(i.key == m.group(1) for i in rmap.insights)
            and m.group(2) in INSIGHT_PLACEHOLDERS
        )
    return False


class SentenceSpec(BaseModel):
    """One sentence of a narrative or insight. ``default`` renders when every placeholder has a
    value; ``zero`` replaces it when the trigger placeholder (the first one in ``default``
    unless ``trigger`` says otherwise) is zero. Without a ``zero`` variant a zero-triggered
    sentence is omitted: a sentence made of zeros is never emitted."""

    default: str
    zero: str | None = None
    one: str | None = Field(
        default=None, description="variant when the trigger is exactly one (singular phrasing)"
    )
    trigger: str | list[str] | None = None
    requires: list[str] = Field(
        default_factory=list,
        description="placeholders that must be present and non-zero for the sentence to render",
    )
    unless: list[str] = Field(
        default_factory=list,
        description="placeholders that must be absent or zero for the sentence to render",
    )

    @property
    def trigger_keys(self) -> list[str]:
        """The placeholders whose all-zero state selects the zero variant."""
        if isinstance(self.trigger, list):
            return self.trigger
        if self.trigger:
            return [self.trigger]
        return placeholder_names(self.default)[:1]

    def templates(self) -> list[str]:
        return [self.default] + [t for t in (self.zero, self.one) if t]

    def placeholders_used(self) -> list[str]:
        out: list[str] = []
        for t in self.templates():
            out.extend(placeholder_names(t))
        return out + self.trigger_keys + list(self.requires) + list(self.unless)


def sentences_of(raw: Any) -> list[SentenceSpec]:
    """Normalise a string, an object or a list of either into sentence specs."""
    if raw is None:
        return []
    if isinstance(raw, str):
        return [SentenceSpec(default=raw)]
    if isinstance(raw, dict):
        return [SentenceSpec.model_validate(raw)]
    if isinstance(raw, list):
        out: list[SentenceSpec] = []
        for item in raw:
            out.extend(sentences_of(item))
        return out
    return [SentenceSpec(default=str(raw))]


class IncludeWhen(BaseModel):
    column: str
    equals: Any


class EntitySpec(BaseModel):
    sheet: str
    rows: tuple[int, int] | None = None
    keyColumn: str
    labelColumn: str | None = None
    includeWhen: IncludeWhen | None = None
    universe: IncludeWhen | None = Field(
        default=None,
        description="a flag column marking rows inside the rated universe; rows outside it "
        "stay in the table but are counted separately (includeWhen drops them instead)",
    )
    asOf: str | None = Field(default=None, description="'Sheet!A1' holding the as-of date")


class DimensionSpec(BaseModel):
    sheet: str | None = None
    column: str
    label: str
    keyColumn: str | None = None
    split: str | None = Field(
        default=None,
        description="separator between the members of one cell (a management team); with "
        "attribute 'team' the cell stays one group rendered as 'the team of A, B and C', with "
        "'person' each member is a group of its own",
    )
    attribute: Literal["team", "person"] = "team"

    def members(self, value: str) -> list[str]:
        if not self.split:
            return [value]
        return [p.strip() for p in value.split(self.split) if p.strip()] or [value]

    def joined(self, value: str) -> str:
        """'A, B and C' for a multi-member cell, the value itself otherwise."""
        parts = self.members(value)
        if len(parts) < 2:
            return parts[0] if parts else value
        return ", ".join(parts[:-1]) + " and " + parts[-1]

    def phrase(self, value: str) -> str:
        """The label a sentence uses: 'the team of A, B and C' for a multi-member cell."""
        parts = self.members(value)
        return f"the team of {self.joined(value)}" if len(parts) >= 2 else value


class MeasureSpec(BaseModel):
    key: str
    label: str
    role: Role
    sheet: str | None = None
    column: str
    keyColumn: str | None = None
    format: str = "number"  # number | integer | percent | general | inr_crore
    unit: str | None = None
    higherIsBetter: bool = True
    primary: bool = False
    decimals: int | None = None


class PhaseGroup(BaseModel):
    key: str
    label: str
    columns: list[str]


class PhaseSpec(BaseModel):
    sheet: str
    keyColumn: str
    headerRow: int
    groups: list[PhaseGroup]
    unit: str | None = None


class PeriodSpec(BaseModel):
    sheet: str
    keyColumn: str
    headerRow: int
    columns: list[str]
    unit: str | None = None


class CategoryStatsSpec(BaseModel):
    sheet: str
    keyColumn: str
    rows: tuple[int, int] | None = None
    headerRow: int
    columns: list[str]
    unit: str | None = None


class QuartileRule(BaseModel):
    unrankedBelow: int = 4
    unrankedValue: str = "--"


class Predicate(BaseModel):
    """A test on a measure (numbers) or a dimension (text). ``value`` may be a literal or
    ``"$name"``, a constant declared under ``research.constants`` and read from a cell."""

    measure: str | None = None
    dimension: str | None = None
    op: Op = "eq"
    value: Any = None
    fraction: float = Field(
        default=0.25, gt=0, le=1, description="share for top/bottom: 0.25 = quarter, 0.5 = half"
    )


class SortSpec(BaseModel):
    measure: str
    dir: Literal["asc", "desc"] = "asc"


class GroupBySpec(BaseModel):
    dimension: str
    aggregate: Literal["count_where", "hit_rate", "mean", "dispersion"] = "count_where"
    where: list[Predicate] = Field(default_factory=list)
    measure: str | None = None
    minMembers: int = 1


class AggregateSpec(BaseModel):
    measure: str
    fn: Literal["sum"] = "sum"


class AcrossSpec(BaseModel):
    where: list[Predicate]
    minVersions: int | None = None


class InsightSpec(BaseModel):
    key: str
    section: str
    eyebrow: str
    title: str
    mode: Literal["filter", "groupBy", "acrossVersions", "aggregate"] = "filter"
    where: list[Predicate] = Field(default_factory=list)
    sort: SortSpec | None = None
    limit: int = 5
    show: list[str] = Field(default_factory=list)
    sentence: Any  # a string or {default, zero}; see SentenceSpec
    groupBy: GroupBySpec | None = None
    aggregate: AggregateSpec | None = None
    across: AcrossSpec | None = None
    drill: dict[str, Any] = Field(default_factory=dict)
    minGroupCount: int | None = Field(
        default=None,
        description="a category (filter) or group (groupBy) counts only with at least this many "
        "rated members; defaults to the map's minGroupCount for groupBy insights",
    )

    @property
    def sentences(self) -> list[SentenceSpec]:
        return sentences_of(self.sentence)


class ConstantSpec(BaseModel):
    cell: str = Field(description="'Sheet!A1' holding the value")
    parse: Literal["number", "firstDate", "lastDate"] = "number"


class CoverageSpec(BaseModel):
    """Why funds are unrated: an inception-date measure against a constant (the earliest phase
    start). Younger funds are too young to rate; older ones with '--' are a data gap."""

    measure: str
    constant: str


class KpiSpec(BaseModel):
    """A dashboard tile: label, a value template and a note template over the executive
    vocabulary (including {ins.<key>.count}), and a tone token (q1, accent, violet, positive)."""

    label: str
    value: str
    note: str | None = None
    tone: Literal["q1", "q2", "accent", "violet", "positive", "warning"] = "accent"
    requires: list[str] = Field(default_factory=list)
    keepZero: bool = Field(
        default=False, description="show the tile even when its value is zero (default: omit)"
    )


class Finding(BaseModel):
    title: str
    detail: str = ""
    status: Literal["fixed", "open"] = "open"


class ResearchMap(BaseModel):
    entity: EntitySpec
    dimensions: dict[str, DimensionSpec] = Field(default_factory=dict)
    measures: list[MeasureSpec]
    phases: PhaseSpec | None = None
    periods: PeriodSpec | None = None
    categoryStats: CategoryStatsSpec | None = None
    quartileRule: QuartileRule = Field(default_factory=QuartileRule)
    minGroupCount: int = Field(
        default=10,
        description="superlatives over groups (AMC, manager, category) ignore groups with fewer "
        "rated members than this, so tiny or exotic groups never top a league table",
    )
    constants: dict[str, ConstantSpec] = Field(default_factory=dict)
    coverage: CoverageSpec | None = None
    sectionLabels: dict[str, str] = Field(default_factory=dict)
    kpis: list[KpiSpec] = Field(default_factory=list)
    insights: list[InsightSpec] = Field(default_factory=list)
    narratives: dict[str, Any] = Field(default_factory=dict)
    footer: str | None = None
    findings: list[Finding] = Field(default_factory=list)

    def measure(self, key: str) -> MeasureSpec | None:
        return next((m for m in self.measures if m.key == key), None)

    def narrative(self, name: str) -> list[SentenceSpec]:
        return sentences_of(self.narratives.get(name))

    def primary(self, role: Role) -> MeasureSpec | None:
        return next((m for m in self.measures if m.role == role and m.primary), None) or next(
            (m for m in self.measures if m.role == role), None
        )


# ---- parsing -----------------------------------------------------------------------------


def parse_map(raw: Any) -> tuple[ResearchMap | None, list[str]]:
    """Validate the raw ``research`` section; problems are readable, never exceptions."""
    if raw is None:
        return None, ["research section is absent from dashboard.config.json"]
    if not isinstance(raw, dict):
        return None, ["research section must be an object"]
    try:
        rmap = ResearchMap.model_validate(raw)
    except ValidationError as exc:
        return None, [
            f"research.{'.'.join(str(p) for p in e['loc'])}: {e['msg']}" for e in exc.errors()
        ]
    problems: list[str] = []
    keys = [m.key for m in rmap.measures]
    for dup in {k for k in keys if keys.count(k) > 1}:
        problems.append(f"measures: key '{dup}' is declared more than once")
    for role in ("score", "rank", "quartile"):
        if rmap.primary(role) is None:  # type: ignore[arg-type]
            problems.append(f"measures: no measure with role '{role}' (a primary one is needed)")
    if "category" not in rmap.dimensions:
        problems.append("dimensions: a 'category' dimension is required")
    if rmap.coverage:
        if rmap.measure(rmap.coverage.measure) is None:
            problems.append(f"coverage: measure '{rmap.coverage.measure}' is not declared")
        if rmap.coverage.constant not in rmap.constants:
            problems.append(f"coverage: constant '{rmap.coverage.constant}' is not declared")
    for i, ins in enumerate(rmap.insights):
        problems.extend(_check_insight(rmap, i, ins))
    for name, raw in rmap.narratives.items():
        if name not in NARRATIVE_PLACEHOLDERS:
            problems.append(
                f"narratives.{name}: unknown narrative (expected one of {sorted(NARRATIVE_PLACEHOLDERS)})"
            )
            continue
        try:
            specs = sentences_of(raw)
        except ValidationError as exc:
            problems.append(f"narratives.{name}: {exc.errors()[0]['msg']}")
            continue
        for spec in specs:
            for ph in dict.fromkeys(spec.placeholders_used()):
                if not narrative_placeholder_ok(name, ph, rmap):
                    problems.append(f"narratives.{name}: unknown placeholder {{{ph}}}")
    for i, kpi in enumerate(rmap.kpis):
        for ph in dict.fromkeys(
            placeholder_names(kpi.value) + placeholder_names(kpi.note or "") + list(kpi.requires)
        ):
            if not narrative_placeholder_ok("executive", ph, rmap):
                problems.append(f"kpis[{i}] '{kpi.label}': unknown placeholder {{{ph}}}")
    return rmap, problems


def _check_insight(rmap: ResearchMap, i: int, ins: InsightSpec) -> list[str]:
    where = f"insights[{i}] '{ins.key}'"
    problems: list[str] = []
    preds = (
        list(ins.where)
        + (list(ins.groupBy.where) if ins.groupBy else [])
        + (list(ins.across.where) if ins.across else [])
    )
    for p in preds:
        if (p.measure is None) == (p.dimension is None):
            problems.append(f"{where}: a predicate needs exactly one of measure or dimension")
        if p.dimension is not None and p.dimension not in rmap.dimensions:
            problems.append(f"{where}: dimension '{p.dimension}' is not declared")
        if (
            isinstance(p.value, str)
            and p.value.startswith("$")
            and p.value[1:] not in rmap.constants
        ):
            problems.append(f"{where}: constant '{p.value}' is not declared under constants")
    refs = [p.measure for p in ins.where if p.measure]
    if ins.sort:
        refs.append(ins.sort.measure)
    if ins.groupBy:
        refs += [p.measure for p in ins.groupBy.where if p.measure]
        if ins.groupBy.measure:
            refs.append(ins.groupBy.measure)
        if ins.groupBy.dimension not in rmap.dimensions:
            problems.append(f"{where}: groupBy dimension '{ins.groupBy.dimension}' is not declared")
        if ins.groupBy.aggregate in ("mean", "dispersion") and not ins.groupBy.measure:
            problems.append(f"{where}: groupBy.{ins.groupBy.aggregate} needs a measure")
    if ins.aggregate:
        refs.append(ins.aggregate.measure)
    if ins.across:
        refs += [p.measure for p in ins.across.where if p.measure]
        primary = {m.key for r in ("score", "rank", "quartile") if (m := rmap.primary(r))}  # type: ignore[arg-type]
        for p in ins.across.where:
            if p.measure and p.measure not in primary and rmap.measure(p.measure) is not None:
                problems.append(
                    f"{where}: across.where can only test the primary score, rank or quartile "
                    f"(version snapshots hold those), not '{p.measure}'"
                )
    if ins.mode == "groupBy" and not ins.groupBy:
        problems.append(f"{where}: mode groupBy needs a groupBy block")
    if ins.mode == "aggregate" and not ins.aggregate:
        problems.append(f"{where}: mode aggregate needs an aggregate block")
    if ins.mode == "acrossVersions" and not ins.across:
        problems.append(f"{where}: mode acrossVersions needs an across block")
    for ref in refs:
        if rmap.measure(ref) is None:
            hint = get_close_matches(ref, [m.key for m in rmap.measures], n=1)
            problems.append(
                f"{where}: measure '{ref}' is not declared"
                + (f"; did you mean '{hint[0]}'?" if hint else "")
            )
    try:
        specs = ins.sentences
    except ValidationError as exc:
        problems.append(f"{where}: sentence: {exc.errors()[0]['msg']}")
        return problems
    if not specs:
        problems.append(f"{where}: sentence is empty")
    for spec in specs:
        for ph in dict.fromkeys(spec.placeholders_used()):
            if ph not in INSIGHT_PLACEHOLDERS:
                problems.append(f"{where}: unknown sentence placeholder {{{ph}}}")
    return problems


# ---- resolution against a version ------------------------------------------------------------


class ResolvedRef(BaseModel):
    sheet: str
    column: str
    col: int
    label: str | None = None  # header text found for the column, when any


class ResolvedMap(BaseModel):
    """The map with every reference checked against one version. ``problems`` lists what did
    not resolve; ``measures`` keeps only the resolved ones (the rest are named in problems)."""

    map: ResearchMap
    problems: list[str] = Field(default_factory=list)
    entity_rows: tuple[int, int]
    entity_key_col: int
    entity_label_col: int | None
    dimensions: dict[str, ResolvedRef] = Field(default_factory=dict)
    measures: dict[str, ResolvedRef] = Field(default_factory=dict)
    unresolved_measures: list[str] = Field(default_factory=list)
    phase_labels: dict[str, dict[str, str]] = Field(default_factory=dict)  # group -> col -> label
    period_labels: dict[str, str] = Field(default_factory=dict)
    category_stat_labels: dict[str, str] = Field(default_factory=dict)
    phases_ok: bool = False
    periods_ok: bool = False
    category_stats_ok: bool = False
    constant_cells: dict[str, tuple[str, int, int]] = Field(
        default_factory=dict
    )  # name -> sheet,row,col

    @property
    def configured(self) -> bool:
        return True

    def measure_specs(self) -> list[MeasureSpec]:
        return [m for m in self.map.measures if m.key in self.measures]


SheetLookup = Callable[[str], Any]  # name -> SheetIndex (app.storage.views)


def _col(letter: str) -> int | None:
    try:
        return column_index_from_string(letter.strip().upper())
    except (ValueError, AttributeError):
        return None


def resolve_map(
    rmap: ResearchMap, model: WorkbookLogicModel, sheet_index: SheetLookup
) -> ResolvedMap | list[str]:
    """Resolve every reference against ``model`` (sheets in scope) and the raw sheets.

    Returns the ResolvedMap (with non-fatal problems inside) or, when the entity itself does not
    resolve, the list of fatal problems."""
    in_scope = {s.name for s in model.sheets if s.in_scope}
    problems: list[str] = []

    def sheet_problem(where: str, name: str) -> str:
        hint = get_close_matches(name, sorted(in_scope), n=1)
        return f"{where}: sheet '{name}' is not in scope" + (
            f"; did you mean '{hint[0]}'?" if hint else ""
        )

    def check_col(where: str, name: str, letter: str) -> int | None:
        idx = sheet_index(name)
        c = _col(letter)
        if c is None:
            problems.append(f"{where}: '{letter}' is not a column letter")
            return None
        if idx.used is None or c > idx.used.c2:
            problems.append(
                f"{where}: column {letter} is outside sheet '{name}' (used range {idx.used.to_a1() if idx.used else 'empty'})"
            )
            return None
        return c

    ent = rmap.entity
    if ent.sheet not in in_scope:
        return [sheet_problem("entity", ent.sheet)]
    eidx = sheet_index(ent.sheet)
    if eidx.used is None:
        return [f"entity: sheet '{ent.sheet}' is empty"]
    key_col = check_col("entity.keyColumn", ent.sheet, ent.keyColumn)
    if key_col is None:
        return problems
    rows = ent.rows or (eidx.used.r1, eidx.used.r2)
    if rows[0] < eidx.used.r1 or rows[1] > eidx.used.r2 or rows[0] > rows[1]:
        return [
            f"entity.rows: {list(rows)} is outside sheet '{ent.sheet}' (used range {eidx.used.to_a1()})"
        ]
    texts = sum(1 for r in range(rows[0], rows[1] + 1) if eidx.any_text(r, key_col))
    span = rows[1] - rows[0] + 1
    if texts < 0.8 * span:
        return [
            f"entity.keyColumn: column {ent.keyColumn} holds text on only {texts} of {span} rows; "
            "it does not look like an identity column"
        ]
    label_col = (
        check_col("entity.labelColumn", ent.sheet, ent.labelColumn) if ent.labelColumn else None
    )
    if ent.includeWhen:
        check_col("entity.includeWhen.column", ent.sheet, ent.includeWhen.column)
    if ent.universe:
        check_col("entity.universe.column", ent.sheet, ent.universe.column)

    resolved = ResolvedMap(
        map=rmap, entity_rows=rows, entity_key_col=key_col, entity_label_col=label_col
    )
    for dkey, d in rmap.dimensions.items():
        name = d.sheet or ent.sheet
        if name not in in_scope:
            problems.append(sheet_problem(f"dimensions.{dkey}", name))
            continue
        c = check_col(f"dimensions.{dkey}", name, d.column)
        if c is not None:
            resolved.dimensions[dkey] = ResolvedRef(sheet=name, column=d.column.upper(), col=c)
    for i, m in enumerate(rmap.measures):
        where = f"measures[{i}] '{m.key}'"
        name = m.sheet or ent.sheet
        if name not in in_scope:
            problems.append(sheet_problem(where, name))
            resolved.unresolved_measures.append(m.key)
            continue
        c = check_col(where, name, m.column)
        if c is None:
            resolved.unresolved_measures.append(m.key)
            continue
        if name != ent.sheet and not m.keyColumn:
            problems.append(f"{where}: measures on another sheet need a keyColumn to join on")
            resolved.unresolved_measures.append(m.key)
            continue
        if m.keyColumn and check_col(f"{where}.keyColumn", name, m.keyColumn) is None:
            resolved.unresolved_measures.append(m.key)
            continue
        resolved.measures[m.key] = ResolvedRef(sheet=name, column=m.column.upper(), col=c)

    def header_labels(
        where: str, name: str, header_row: int, columns: list[str]
    ) -> dict[str, str] | None:
        if name not in in_scope:
            problems.append(sheet_problem(where, name))
            return None
        idx = sheet_index(name)
        labels: dict[str, str] = {}
        ok = True
        for letter in columns:
            c = check_col(where, name, letter)
            if c is None:
                ok = False
                continue
            text = idx.any_text(header_row, c)
            if not text:
                problems.append(f"{where}: header row {header_row} has no text in column {letter}")
                ok = False
                continue
            labels[letter.upper()] = text
        return labels if ok else None

    if rmap.phases:
        p = rmap.phases
        ok = (
            check_col("phases.keyColumn", p.sheet, p.keyColumn) is not None
            if p.sheet in in_scope
            else False
        )
        for g in p.groups:
            labels = header_labels(f"phases.{g.key}", p.sheet, p.headerRow, g.columns)
            if labels is None:
                ok = False
            else:
                resolved.phase_labels[g.key] = labels
        resolved.phases_ok = ok and bool(resolved.phase_labels)
    if rmap.periods:
        p = rmap.periods
        labels = header_labels("periods", p.sheet, p.headerRow, p.columns)
        ok = (
            p.sheet in in_scope and check_col("periods.keyColumn", p.sheet, p.keyColumn) is not None
        )
        if labels is not None and ok:
            resolved.period_labels = labels
            resolved.periods_ok = True
    if rmap.categoryStats:
        cs = rmap.categoryStats
        labels = header_labels("categoryStats", cs.sheet, cs.headerRow, cs.columns)
        ok = (
            cs.sheet in in_scope
            and check_col("categoryStats.keyColumn", cs.sheet, cs.keyColumn) is not None
        )
        if labels is not None and ok:
            resolved.category_stat_labels = labels
            resolved.category_stats_ok = True
    for name, spec in rmap.constants.items():
        if "!" not in spec.cell:
            problems.append(f"constants.{name}: '{spec.cell}' is not 'Sheet!A1'")
            continue
        sheet_name, addr = spec.cell.rsplit("!", 1)
        sheet_name = sheet_name.strip("'")
        if sheet_name not in in_scope:
            problems.append(sheet_problem(f"constants.{name}", sheet_name))
            continue
        try:
            from app.model.formula.refs import parse_a1_cell

            r, c = parse_a1_cell(addr)
        except Exception:  # noqa: BLE001
            problems.append(f"constants.{name}: '{addr}' is not a cell address")
            continue
        resolved.constant_cells[name] = (sheet_name, r, c)
    for role in ("score", "rank", "quartile"):
        prim = rmap.primary(role)  # type: ignore[arg-type]
        if prim is not None and prim.key not in resolved.measures:
            problems.append(
                f"primary {role} measure '{prim.key}' did not resolve; headline views need it"
            )
    resolved.problems = problems
    return resolved


def column_text(sheet: str, letter: str) -> str:
    return f"{sheet}!{letter}"


def letter_of(col: int) -> str:
    return get_column_letter(col)
