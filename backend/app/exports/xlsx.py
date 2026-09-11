"""Excel export: a cover sheet plus one sheet per view, streamed with openpyxl's write-only mode.

In-place views (``first_row`` set) land cell for cell at their original addresses, with the
original number formats, so an exported sheet reads exactly like the source sheet with computed
values where the formulas were. Errors are written as Excel error cells, not text.
"""

from __future__ import annotations

from collections.abc import Iterable
from typing import IO, Any

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from app.exports.descriptor import ExportView, Provenance, slug

ERROR_TEXTS = {"#NULL!", "#DIV/0!", "#VALUE!", "#REF!", "#NAME?", "#NUM!", "#N/A"}


def _cell(ws, value: Any, vtype: str | None, fmt: str | None):
    if value is None:
        return None
    if isinstance(value, bool):
        return value
    if vtype == "error" or (isinstance(value, str) and value in ERROR_TEXTS):
        c = WriteOnlyCell(ws, value=str(value))
        c.data_type = "e"
        return c
    if isinstance(value, int | float) and fmt and fmt != "General":
        c = WriteOnlyCell(ws, value=float(value))
        c.number_format = fmt
        return c
    if isinstance(value, float) and value.is_integer() and abs(value) < 1e15:
        return int(value)
    return value


def _cover(wb: Workbook, provenance: Provenance, views: Iterable[ExportView]) -> None:
    ws = wb.create_sheet("Cover")
    bold = Font(bold=True)
    title = WriteOnlyCell(ws, value=provenance.workbook or provenance.filename)
    title.font = Font(bold=True, size=14)
    ws.append([title])
    ws.append([])
    for key, value in provenance.lines():
        k = WriteOnlyCell(ws, value=key)
        k.font = bold
        ws.append([k, value])
    ws.append([])
    head = WriteOnlyCell(ws, value="Sheets in this file")
    head.font = bold
    ws.append([head, "Rows", "Columns", "Scope"])
    for v in views:
        ws.append([v.title, len(v.rows), len(v.columns), v.subtitle or ""])


def _sheet_name(view: ExportView, taken: set[str]) -> str:
    base = slug(view.title, 31) or view.id
    name = base
    n = 2
    while name in taken:
        suffix = f"~{n}"
        name = base[: 31 - len(suffix)] + suffix
        n += 1
    taken.add(name)
    return name


def write_workbook(views: list[ExportView], provenance: Provenance, target: IO[bytes]) -> None:
    wb = Workbook(write_only=True)
    _cover(wb, provenance, views)
    taken = {"Cover"}
    for view in views:
        ws = wb.create_sheet(_sheet_name(view, taken))
        if view.first_row is not None:
            # In place: pad to the original row and column so addresses match the source.
            for _ in range(1, view.first_row):
                ws.append([])
            pad = [None] * ((view.first_col or 1) - 1)
            for r, values in enumerate(view.rows):
                types = view.types[r] if view.types else [None] * len(values)
                fmts = view.formats[r] if view.formats else [None] * len(values)
                ws.append(
                    pad + [_cell(ws, v, t, f) for v, t, f in zip(values, types, fmts, strict=True)]
                )
            continue
        header = [WriteOnlyCell(ws, value=view.row_label_header)] + [
            WriteOnlyCell(ws, value=c.label) for c in view.columns
        ]
        for h in header:
            h.font = Font(bold=True)
        ws.append(header)
        if view.columns and view.columns[0].kind != "field":
            ws.append([""] + [c.key for c in view.columns])
        for r, values in enumerate(view.rows):
            types = view.types[r] if view.types else [None] * len(values)
            fmts = view.formats[r] if view.formats else [None] * len(values)
            label = view.row_labels[r] if r < len(view.row_labels) else None
            ws.append(
                [label] + [_cell(ws, v, t, f) for v, t, f in zip(values, types, fmts, strict=True)]
            )
    wb.save(target)


def column_letters(n: int) -> list[str]:
    return [get_column_letter(i) for i in range(1, n + 1)]
